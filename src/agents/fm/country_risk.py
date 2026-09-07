"""
Country risk premiums and sovereign default spreads, from Damodaran's table.

Aswath Damodaran (NYU Stern) publishes, twice a year, a country-by-country
table of Moody's ratings, the default spread implied by each rating, and the
country risk premium — that spread scaled by how much more volatile equities
are than government bonds (x1.52 in the January 2026 update). It is the
reference most practitioners use, and it is what this module reads.

The premiums used to be a hand-typed table of "approximate, mid-2025" figures
with every developed market at zero. Damodaran's are not: the United States
carries 0.23% since its May 2025 downgrade to Aa1, France and the UK 0.78%,
Japan 0.91%. They enter the cost of equity as Ke = Rf + beta x (ERP + CRP).

The same table supplies the sovereign default spread that the risk-free build
subtracts from a government bond yield (see assumption_grounding): a bond
rated Baa3 is not default-free, and the default risk it carries is exactly
what the country premium prices — leaving it in Rf as well would charge it
twice.

The spreadsheet is fetched at run time (cached seven days); the snapshot in
country_risk_snapshot.py serves when it cannot be. Both are dated, the date is
printed, and CRP_<COUNTRY> overrides a premium without a deploy.
"""

from __future__ import annotations

import io
import os
import threading
from typing import Any, Dict, Optional, Tuple

from . import netcache
from . import country_risk_snapshot as _snap

_URL = "https://pages.stern.nyu.edu/~adamodar/pc/datasets/ctryprem.xlsx"
_TIMEOUT = 20.0
_CACHE_TTL = 7 * 86400
_MIN_COUNTRIES = 100          # a parse that finds fewer has read the wrong sheet

# Yahoo's country strings, and common variants, mapped to Damodaran's names.
_ALIASES = {
    "south korea": "Korea", "korea, republic of": "Korea", "republic of korea": "Korea",
    "czechia": "Czech Republic", "russian federation": "Russia", "viet nam": "Vietnam",
    "uae": "United Arab Emirates", "macau": "Macao", "türkiye": "Turkey", "turkiye": "Turkey",
    "ivory coast": "Côte d'Ivoire", "cote d'ivoire": "Côte d'Ivoire",
    "usa": "United States", "united states of america": "United States",
    "uk": "United Kingdom", "great britain": "United Kingdom",
    "hong kong sar": "Hong Kong", "hong kong sar china": "Hong Kong",
    "taiwan, province of china": "Taiwan", "people's republic of china": "China", "prc": "China",
    "north macedonia": "Macedonia", "eswatini": "Swaziland", "slovak republic": "Slovakia",
    "bosnia": "Bosnia and Herzegovina", "cayman": "Cayman Islands",
    "trinidad & tobago": "Trinidad and Tobago", "st. vincent": "St. Vincent & the Grenadines",
    "jersey": "Jersey (States of)", "guernsey": "Guernsey (States of)", "andorra": "Andorra (Principality of)",
}

# The sovereign whose bond stands behind each currency.
_CURRENCY_SOVEREIGN = {
    "USD": "United States", "EUR": "Germany", "GBP": "United Kingdom", "JPY": "Japan",
    "CHF": "Switzerland", "INR": "India", "KRW": "Korea", "CNY": "China", "HKD": "Hong Kong",
    "TWD": "Taiwan", "SGD": "Singapore", "AUD": "Australia", "CAD": "Canada", "BRL": "Brazil",
    "MXN": "Mexico", "ZAR": "South Africa", "SEK": "Sweden", "NOK": "Norway", "DKK": "Denmark",
    "NZD": "New Zealand", "PLN": "Poland", "ILS": "Israel", "HUF": "Hungary",
    "CZK": "Czech Republic", "CLP": "Chile", "IDR": "Indonesia", "TRY": "Turkey",
    "SAR": "Saudi Arabia", "AED": "United Arab Emirates", "THB": "Thailand", "MYR": "Malaysia",
    "PHP": "Philippines", "VND": "Vietnam", "ARS": "Argentina", "EGP": "Egypt", "NGN": "Nigeria",
    "PKR": "Pakistan", "COP": "Colombia", "PEN": "Peru", "RUB": "Russia", "QAR": "Qatar",
    "KWD": "Kuwait", "RON": "Romania", "BGN": "Bulgaria", "ISK": "Iceland", "KZT": "Kazakhstan",
    "BDT": "Bangladesh", "LKR": "Sri Lanka", "KES": "Kenya", "MAD": "Morocco", "UAH": "Ukraine",
}

_memo: Dict[str, Any] = {}
_lock = threading.Lock()


def _reset_memo() -> None:
    with _lock:
        _memo.clear()


def _update_label(as_of: Optional[str]) -> str:
    """'2026-01-01' -> 'January 2026 update'."""
    import datetime as dt
    try:
        return dt.date.fromisoformat(str(as_of)[:10]).strftime("%B %Y") + " update"
    except (ValueError, TypeError):
        return f"{as_of} update"


def parse_workbook(data: bytes) -> Dict[str, Any]:
    """
    The 'ERPs by country' sheet: a 'Date of update' cell, the mature-market
    ERP, then a header row (Country / Moody's rating / Rating-based Default
    Spread / Country Risk Premium ...) and one row per country. Frontier
    markets and a few malformed rows have numbers where the rating should be;
    they are skipped, not guessed.
    """
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if "ERPs by country" not in wb.sheetnames:
        raise ValueError("no 'ERPs by country' sheet")
    ws = wb["ERPs by country"]

    as_of = None
    mature = None
    countries: Dict[str, Dict[str, Any]] = {}
    cols = None
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        first = row[0]
        if not isinstance(first, str):
            continue
        low = first.strip().lower()
        if cols is None:
            if low.startswith("date of update") and len(row) > 1:
                v = row[1]
                as_of = v.date().isoformat() if hasattr(v, "date") else str(v)[:10]
            elif "mature equity market" in low:
                mature = next((float(c) for c in row[1:] if isinstance(c, (int, float)) and 0 < c < 0.2), None)
            elif first.strip() == "Country":
                hdr = [str(c).strip().lower() if isinstance(c, str) else "" for c in row]
                try:
                    cols = (next(i for i, h in enumerate(hdr) if h.startswith("moody")),
                            next(i for i, h in enumerate(hdr) if h.startswith("rating-based default spread")),
                            next(i for i, h in enumerate(hdr) if h == "country risk premium"))
                except StopIteration:
                    raise ValueError("header row without the expected columns")
            continue
        i_rating, i_ds, i_crp = cols
        if max(cols) >= len(row):
            continue
        rating, ds, crp = row[i_rating], row[i_ds], row[i_crp]
        if not isinstance(rating, str) or isinstance(ds, bool) or isinstance(crp, bool):
            continue
        if not isinstance(ds, (int, float)) or not isinstance(crp, (int, float)):
            continue
        if not (0.0 <= ds <= 0.5 and 0.0 <= crp <= 0.5):
            continue
        countries[first.strip()] = {"rating": rating.strip(), "default_spread": float(ds), "crp": float(crp)}

    if len(countries) < _MIN_COUNTRIES or not as_of:
        raise ValueError(f"unexpected layout: {len(countries)} countries, as_of={as_of}")
    return {"as_of": as_of, "mature_erp": mature, "countries": countries}


def _snapshot_table() -> Dict[str, Any]:
    return {
        "as_of": _snap.AS_OF,
        "mature_erp": _snap.MATURE_ERP,
        "countries": {k: {"rating": r, "default_spread": d, "crp": c} for k, (r, d, c) in _snap.COUNTRIES.items()},
        "source": f"Damodaran {_update_label(_snap.AS_OF)}, snapshot",
    }


def _valid_table(table: object) -> bool:
    """A disk copy is untrusted input; anything not shaped like the table is a miss."""
    if not isinstance(table, dict) or not isinstance(table.get("source"), str):
        return False
    countries = table.get("countries")
    if not isinstance(countries, dict) or len(countries) < _MIN_COUNTRIES:
        return False
    for entry in countries.values():
        if not isinstance(entry, dict) or not isinstance(entry.get("rating"), str):
            return False
        for k in ("default_spread", "crp"):
            v = entry.get(k)
            if isinstance(v, bool) or not isinstance(v, (int, float)) or not (0.0 <= v <= 0.5):
                return False
    return True


def load_table() -> Dict[str, Any]:
    """The current table: memo -> disk cache (7 days) -> fetch -> older disk copy -> snapshot."""
    with _lock:
        if _memo.get("table"):
            return _memo["table"]

    table = netcache.cache_get("ctryprem", _CACHE_TTL)
    if not _valid_table(table):
        table = None
        try:
            table = parse_workbook(netcache.http_get(_URL, _TIMEOUT))
            table["source"] = f"Damodaran {_update_label(table['as_of'])}"
            netcache.cache_put("ctryprem", table)
        except Exception:
            table = None
    if table is None:
        # The fetch failed. A copy from an earlier run — the July update
        # fetched eight days ago — is newer than the embedded snapshot.
        older = netcache.cache_get("ctryprem", 400 * 86400)
        if _valid_table(older):
            table = dict(older)
            if "cached" not in table["source"]:
                table["source"] = f"{table['source']}, cached"
    if table is None:
        table = _snapshot_table()

    with _lock:
        _memo["table"] = table
    return table


def country_entry(country: Optional[str]) -> Optional[Dict[str, Any]]:
    """The table row for ``country`` under any of its usual names, or None."""
    name = (country or "").strip()
    if not name:
        return None
    countries = load_table()["countries"]
    hit = countries.get(name)
    if hit is None:
        alias = _ALIASES.get(name.lower())
        if alias:
            hit = countries.get(alias)
            name = alias
    if hit is None:
        low = name.lower()
        for k, v in countries.items():
            # "Jersey" -> "Jersey (States of)", "Congo" -> "Congo (Republic of)"
            if k.lower() == low or k.lower().startswith(low + " ("):
                hit, name = v, k
                break
    return dict(hit, name=name) if hit else None


def country_risk_premium(country: Optional[str]) -> Tuple[float, str]:
    """
    Premium over the mature-market ERP for ``country``, with its provenance.

    Returns ``(premium, label)``. Unknown countries get no premium and a label
    that says so — the alternative, guessing, is how numbers stop being
    arguable.
    """
    name = (country or "").strip()
    if not name:
        return 0.0, "no country on record — no country premium applied"

    key = "CRP_" + name.upper().replace(" ", "_")
    override = os.getenv(key)
    if override:
        try:
            value = float(override)
            if 0.0 <= value <= 0.30:
                return value, f"{name} {value*100:.1f}% ({key} override)"
        except ValueError:
            pass

    entry = country_entry(name)
    if entry is None:
        return 0.0, f"{name}: no country premium on record — none applied"
    source = load_table()["source"]
    if entry["crp"] <= 0.0:
        return 0.0, f"{entry['name']}: Moody's {entry['rating']}, no country premium ({source})"
    return entry["crp"], f"{entry['name']} {entry['crp']*100:.2f}% ({source}; Moody's {entry['rating']})"


def sovereign_for_currency(currency: Optional[str]) -> Optional[str]:
    """The country whose government bond stands behind ``currency``, if known."""
    from .sovereign_rates import normalise_currency
    return _CURRENCY_SOVEREIGN.get(normalise_currency(currency))


def sovereign_default_spread(currency: Optional[str]) -> Tuple[float, str]:
    """
    The rating-based default spread of the sovereign behind ``currency``.

    Zero, with an empty label, when the currency's sovereign is unknown or
    rated Aaa.
    """
    from .sovereign_rates import normalise_currency
    sovereign = _CURRENCY_SOVEREIGN.get(normalise_currency(currency))
    entry = country_entry(sovereign) if sovereign else None
    if not entry or entry["default_spread"] <= 0.0:
        return 0.0, ""
    return entry["default_spread"], f"{entry['name']} Moody's {entry['rating']}, {load_table()['source']}"
