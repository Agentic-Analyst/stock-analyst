"""Read-only audit of saved valuation artifacts under current publication rules."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, Optional, Sequence

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from src.external_expectations import (
    build_external_expectations,
    implied_discount_rate_for_enterprise_value,
    implied_fcf_path_scale_for_enterprise_value,
    implied_terminal_growth_for_enterprise_value,
)
from src.agents.fm.formula_evaluator import (
    FormulaEvaluator,
    formula_integrity,
    model_integrity,
)
from src.valuation_methodology import normalize_peer_comps_policy
from report_agent import (
    apply_valuation_override,
    enforce_valuation_publication_boundary,
    extract_company_overview,
    extract_projections,
    extract_valuation,
    valuation_override_from_publication_metadata,
)


def _json(path: Path) -> Dict[str, Any]:
    value = json.loads(path.read_text())
    if not isinstance(value, dict):
        raise ValueError(f"Expected an object in {path}")
    return value


def _single(root: Path, pattern: str) -> Path:
    matches = sorted(root.glob(pattern))
    if len(matches) != 1:
        raise ValueError(f"Expected one {pattern} below {root}; found {len(matches)}")
    return matches[0]


def _replay_formula_integrity(root: Path) -> Dict[str, Any]:
    """Evaluate the workbook itself when legacy JSON has no integrity manifest.

    Older artifacts predate ``_vynn.formula_integrity``. Treating a missing
    manifest as success allowed a bank workbook with 127 broken, inapplicable
    industrial-DCF formulas to pass the artifact audit. Replaying the XLSX is
    read-only and proves what a user would actually download.
    """
    workbook_path = _single(root, "models/*.xlsx")
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    results = FormulaEvaluator(workbook).evaluate_all_tabs()
    return formula_integrity(results)


def _replay_model_integrity(root: Path) -> Dict[str, Any]:
    """Rebuild semantic identity checks for an artifact lacking a manifest."""
    workbook_path = _single(root, "models/*.xlsx")
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    results = FormulaEvaluator(workbook).evaluate_all_tabs()
    return model_integrity(results)


def _workbook_headline_label(root: Path) -> Optional[str]:
    """Read the label users actually download, never infer it from sidecar JSON."""
    workbook_path = _single(root, "models/*.xlsx")
    workbook = load_workbook(
        workbook_path, data_only=False, read_only=True,
    )
    if "Summary" not in workbook.sheetnames:
        return None
    value = workbook["Summary"]["A26"].value
    return value if isinstance(value, str) else None


def _withheld_headline_is_explicit(label: Optional[str]) -> bool:
    """A hidden number needs an unmistakable non-publication label.

    Merely avoiding the words ``fair value`` is not enough.  Legacy workbooks
    used labels such as ``Average of Methods (Per-Share)`` and
    ``Bank Valuation Audit Midpoint`` beside a prominent price.  A reasonable
    reader can still mistake either for the product's answer even when current
    policy rejects the point estimate.  Require an explicit withholding phrase
    at the exact downloadable headline cell.
    """
    if not isinstance(label, str):
        return False
    normalized = " ".join(label.casefold().split())
    return any(marker in normalized for marker in (
        "not published",
        "publication withheld",
        "point estimate withheld",
    ))


def _publication_input_from_artifact(
    financial: Dict[str, Any], computed: Dict[str, Any],
) -> Dict[str, Any]:
    """Reconstruct the exact report input needed by the publication boundary.

    In particular, preserve the saved forecast clock.  A rolling NTM model
    must be compared with Street estimates blended to that same clock; dropping
    ``model_inputs.forecast_basis`` silently compares it with raw fiscal-year
    rows and can reverse the current-policy audit decision.
    """
    internal = (computed.get("_vynn") or {}) if isinstance(computed, dict) else {}
    return {
        "company_overview": extract_company_overview(financial),
        "projections": extract_projections(computed),
        "valuation": extract_valuation(computed),
        "model_inputs": internal.get("model_inputs") or {},
    }


def _specialized_refusal_without_model(
    root: Path, financial: Dict[str, Any],
) -> Optional[Dict[str, Any]]:
    """Audit an intentional pre-model refusal as a successful safety outcome.

    REIT, insurer, commodity-cycle, fund, crypto, and unsupported-company
    classifications can stop before a corporate workbook is created.  The
    previous auditor unconditionally required one XLSX and one computed JSON,
    so the safest possible result was reported as an opaque ``ValueError``.
    Only accept the refusal when the methodology explicitly names a specialized
    service, explicitly disallows publication, and *neither* model artifact
    exists.  A partial or stray model still falls through to the strict audit.
    """
    methodology = financial.get("valuation_methodology") or {}
    service = methodology.get("specialized_service")
    publication_allowed = methodology.get("publication_allowed")
    model_files = sorted(root.glob("models/*"))
    if not service or publication_allowed is not False or model_files:
        return None

    company = extract_company_overview(financial)
    external = (
        financial.get("external_expectations")
        or build_external_expectations(financial)
    )
    target = external.get("price_target") or {}
    observations = external.get("analyst_observations") or {}
    reason = methodology.get("reason") or (
        f"Corporate valuation is not applicable; use the {service} service."
    )
    return {
        "status": "passed",
        "run": str(root),
        "artifact_kind": "specialized_refusal_without_corporate_model",
        "ticker": financial.get("ticker") or company.get("ticker"),
        "quote_type": ((financial.get("company_data") or {}).get(
            "basic_info") or {}).get("quote_type"),
        "financial_currency": company.get("currency"),
        "listing_currency": company.get("listing_currency"),
        "current_price": company.get("current_price"),
        "current_price_listing": company.get("current_price_listing"),
        "fx_listing_to_financial": company.get("fx_listing_to_financial"),
        "method": methodology.get("primary_method"),
        "specialized_service": service,
        "current_policy_audit_value": None,
        "current_policy_publishable_value": None,
        "point_estimate_withheld": True,
        "withheld_reason": reason,
        "financial_freshness": financial.get("financial_freshness") or {},
        "street_target": target.get("mean"),
        "street_target_count": target.get("analyst_count"),
        "street_target_gap_vs_market": target.get("return_vs_market"),
        "street_coverage": external.get("coverage"),
        "analyst_observation_count": observations.get("observation_count"),
        "analyst_observation_as_of": observations.get("as_of"),
        "analyst_observations_included_in_intrinsic_value": (
            observations.get("included_in_intrinsic_value")
        ),
        "checks": [
            "PASS specialized methodology refused an inapplicable corporate model"
        ],
    }


def audit_run(root: Path) -> Dict[str, Any]:
    """Reapply current deterministic controls without changing the artifact."""
    financial = _json(_single(root, "financials/financials_annual_modeling_latest.json"))
    specialized = _specialized_refusal_without_model(root, financial)
    if specialized is not None:
        return specialized
    computed = _json(_single(root, "models/*_computed_values.json"))
    external = financial.get("external_expectations") or build_external_expectations(financial)
    peers = ((financial.get("industry_data") or {}).get("peer_comps") or {})
    peer_policy = normalize_peer_comps_policy(peers)
    original = extract_valuation(computed)
    company = extract_company_overview(financial)
    original_summary = original.get("summary") or {}
    original_headline_value = original_summary.get("average_intrinsic")
    internal = (computed.get("_vynn") or {}) if isinstance(computed, dict) else {}
    stored_publication = internal.get("valuation_publication") or {}
    stored_formula_integrity = internal.get("formula_integrity") or {}
    stored_model_integrity = internal.get("model_integrity") or {}
    if stored_formula_integrity:
        integrity = stored_formula_integrity
        integrity_source = "stored_manifest"
    else:
        integrity = _replay_formula_integrity(root)
        integrity_source = "workbook_replay"
    if stored_model_integrity:
        semantic_integrity = stored_model_integrity
        semantic_integrity_source = "stored_manifest"
    else:
        semantic_integrity = _replay_model_integrity(root)
        semantic_integrity_source = "workbook_replay"
    workbook_headline_label = _workbook_headline_label(root)
    data = _publication_input_from_artifact(financial, computed)
    data = apply_valuation_override(
        data, valuation_override_from_publication_metadata(computed)
    )
    data = enforce_valuation_publication_boundary(data, financial)
    valuation = data.get("valuation") or {}
    bank = valuation.get("bank") or {}
    summary = valuation.get("summary") or {}
    reliability = valuation.get("reliability") or {}
    suitability = reliability.get("method_suitability") or {}
    target = external.get("price_target") or {}
    analyst_observations = external.get("analyst_observations") or {}
    current = company.get("current_price")
    dcf_values = ([None, None] if bank else [
        (valuation.get("dcf_perpetual") or {}).get("intrinsic_value_per_share"),
        (valuation.get("dcf_exit") or {}).get("intrinsic_value_per_share"),
    ])
    valid_dcf = [float(value) for value in dcf_values
                 if isinstance(value, (int, float)) and not isinstance(value, bool)
                 and value > 0]
    dcf_midpoint = sum(valid_dcf) / len(valid_dcf) if valid_dcf else None
    model_ev = (valuation.get("dcf_perpetual") or {}).get("enterprise_value")
    market_path_scale = implied_fcf_path_scale_for_enterprise_value(
        (valuation.get("reverse_dcf") or {}).get("market_enterprise_value"),
        model_enterprise_value=model_ev,
    )
    target_path_scale = implied_fcf_path_scale_for_enterprise_value(
        (external.get("valuation_cross_check") or {}).get(
            "target_implied_enterprise_value"
        ),
        model_enterprise_value=model_ev,
    )
    dcf_inputs = valuation.get("dcf_inputs") or {}
    explicit_fcf = dcf_inputs.get("fcf") or []
    model_wacc = dcf_inputs.get("wacc")
    model_terminal_growth = dcf_inputs.get("terminal_growth")
    mid_year_adjustment = dcf_inputs.get("mid_year_adjustment", 0.0)
    market_ev = (valuation.get("reverse_dcf") or {}).get(
        "market_enterprise_value"
    )
    target_ev = (external.get("valuation_cross_check") or {}).get(
        "target_implied_enterprise_value"
    )
    reverse_assumptions = {}
    for prefix, benchmark_ev in (("market", market_ev), ("analyst_target", target_ev)):
        rate = implied_discount_rate_for_enterprise_value(
            benchmark_ev,
            explicit_fcf=explicit_fcf,
            terminal_growth=model_terminal_growth,
            model_wacc=model_wacc,
            mid_year_adjustment=mid_year_adjustment,
        )
        growth = implied_terminal_growth_for_enterprise_value(
            benchmark_ev,
            explicit_fcf=explicit_fcf,
            wacc=model_wacc,
            model_terminal_growth=model_terminal_growth,
            mid_year_adjustment=mid_year_adjustment,
        )
        reverse_assumptions[f"{prefix}_implied_wacc"] = rate.get("implied_wacc")
        reverse_assumptions[f"{prefix}_implied_wacc_vs_model"] = rate.get(
            "implied_wacc_vs_model"
        )
        reverse_assumptions[f"{prefix}_implied_terminal_growth"] = growth.get(
            "implied_terminal_growth"
        )
        reverse_assumptions[
            f"{prefix}_implied_terminal_growth_vs_model"
        ] = growth.get("implied_terminal_growth_vs_model")
    checks = []
    if peer_policy["broad_sector"] and peer_policy["included_in_blended_value"]:
        checks.append("FAIL broad sector peers retain a valuation vote")
    if suitability.get("specialized_service") and not reliability.get("point_estimate_withheld"):
        checks.append("FAIL specialized asset passed corporate publication boundary")
    if reliability.get("point_estimate_withheld") and not reliability.get("withheld_reason"):
        checks.append("FAIL withheld output lacks a reason")
    range_low = reliability.get("range_low")
    range_high = reliability.get("range_high")
    if (range_low is None) != (range_high is None):
        checks.append("FAIL supported valuation range has only one endpoint")
    if range_low is not None and (
            not isinstance(range_low, (int, float)) or isinstance(range_low, bool)
            or not isinstance(range_high, (int, float)) or isinstance(range_high, bool)
            or range_low <= 0 or range_high <= 0 or range_low > range_high):
        checks.append("FAIL supported valuation range is nonpositive or reversed")
    if not reliability.get("point_estimate_withheld") and (
            not isinstance(summary.get("average_intrinsic"), (int, float))
            or isinstance(summary.get("average_intrinsic"), bool)
            or summary.get("average_intrinsic") <= 0):
        checks.append("FAIL publishable output lacks a positive headline value")
    if integrity.get("status") != "ready":
        checks.append("FAIL saved workbook formula integrity was not ready")
    if semantic_integrity.get("status") != "ready":
        checks.append(
            "FAIL saved workbook accounting or valuation identity integrity was not ready"
        )
    if (reliability.get("point_estimate_withheld") and
            not _withheld_headline_is_explicit(workbook_headline_label)):
        checks.append(
            "FAIL downloadable workbook headline does not explicitly say that "
            "the point estimate is not published"
        )
    if (stored_publication.get("publication_allowed") is True and
            reliability.get("point_estimate_withheld")):
        checks.append("FAIL saved artifact publication permission is rejected by current policy")
    if stored_publication.get("status") == "error":
        checks.append("WARN saved artifact publication metadata failed; regenerate before use")
    if analyst_observations.get("included_in_intrinsic_value") is True:
        checks.append("FAIL external analyst observations retain an intrinsic-value vote")
    if (analyst_observations.get("observation_count")
            and not analyst_observations.get("as_of")):
        checks.append("FAIL external analyst observations lack a provider observation date")
    failed = any(str(check).startswith("FAIL ") for check in checks)
    warned = any(str(check).startswith("WARN ") for check in checks)
    if not checks:
        checks.append("PASS current deterministic publication invariants")
    return {
        "status": (
            "failed" if failed else
            "passed_with_warnings" if warned else
            "passed"
        ),
        "run": str(root),
        "ticker": financial.get("ticker") or company.get("ticker"),
        "quote_type": ((financial.get("company_data") or {}).get("basic_info") or {}).get(
            "quote_type"
        ),
        # ``current_price`` and all model values are denominated in the
        # financial-statement currency.  Cross-currency/ADR launch reviews
        # also need the exchange quote and bridge; omitting these made a
        # correct TWD-per-ADR TSM result look like an implausible USD price.
        "financial_currency": company.get("currency"),
        "listing_currency": company.get("listing_currency"),
        "current_price": current,
        "current_price_listing": company.get("current_price_listing"),
        "fx_listing_to_financial": company.get("fx_listing_to_financial"),
        "original_headline_value": original_headline_value,
        # The boundary retains scenario arithmetic for audit. This is not the
        # public value unless point publication is explicitly allowed.
        "current_policy_audit_value": summary.get("average_intrinsic"),
        "current_policy_publishable_value": (
            None if reliability.get("point_estimate_withheld")
            else summary.get("average_intrinsic")
        ),
        "workbook_headline_label": workbook_headline_label,
        "stored_publication_allowed": stored_publication.get("publication_allowed"),
        "stored_publication_status": stored_publication.get("status"),
        "formula_integrity_status": integrity.get("status"),
        "formula_integrity_source": integrity_source,
        "formula_integrity_issue_count": integrity.get("issue_count"),
        "formula_integrity_issues": (integrity.get("issues") or [])[:10],
        "model_integrity_status": semantic_integrity.get("status"),
        "model_integrity_source": semantic_integrity_source,
        "model_integrity_issue_count": semantic_integrity.get("issue_count"),
        "model_integrity_issues": (semantic_integrity.get("issues") or [])[:10],
        "dcf_perpetual": dcf_values[0],
        "dcf_exit": dcf_values[1],
        "dcf_midpoint": dcf_midpoint,
        "dcf_gap_vs_market": dcf_midpoint / current - 1
        if dcf_midpoint is not None and isinstance(current, (int, float)) and current > 0 else None,
        "peer_value": (bank.get("peer_fair_value") if bank
                       else summary.get("comps_intrinsic")),
        "bank_intrinsic_value": bank.get("intrinsic_fair_value") if bank else None,
        "bank_forward_consensus_value": (
            bank.get("forward_consensus_fair_value") if bank else None
        ),
        "bank_peer_value": bank.get("peer_fair_value") if bank else None,
        "peer_policy": peer_policy,
        "street_target": target.get("mean"),
        "street_target_count": target.get("analyst_count"),
        "street_target_gap_vs_market": target.get("return_vs_market"),
        "street_coverage": external.get("coverage"),
        "analyst_observation_count": analyst_observations.get(
            "observation_count"
        ),
        "analyst_observation_as_of": analyst_observations.get("as_of"),
        "analyst_observations_included_in_intrinsic_value": (
            analyst_observations.get("included_in_intrinsic_value")
        ),
        "method": suitability.get("primary_method"),
        "freshness": reliability.get("financial_freshness"),
        "reliability_band": reliability.get("band"),
        "range_low": range_low,
        "range_high": range_high,
        "method_values_for_audit": reliability.get("legs") or {},
        "failed_method_values": reliability.get("failed_legs") or {},
        "point_estimate_withheld": reliability.get("point_estimate_withheld"),
        "withheld_reason": reliability.get("withheld_reason"),
        "reverse_dcf_market_implied_vs_model": (
            (valuation.get("reverse_dcf") or {}).get("market_implied_vs_model")
        ),
        "market_implied_fcf_path_vs_model": market_path_scale.get(
            "implied_fcf_path_vs_model"
        ),
        "analyst_target_implied_fcf_path_vs_model": target_path_scale.get(
            "implied_fcf_path_vs_model"
        ),
        "model_wacc": model_wacc,
        "model_tax_rate": dcf_inputs.get("tax_rate"),
        "model_terminal_growth": model_terminal_growth,
        **reverse_assumptions,
        "checks": checks,
    }


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("runs", nargs="+", type=Path)
    args = parser.parse_args(argv)
    rows = []
    for root in args.runs:
        try:
            rows.append(audit_run(root))
        except Exception as error:
            rows.append({
                "run": str(root), "status": "failed",
                "error_type": type(error).__name__,
            })
    print(json.dumps(rows, indent=2, sort_keys=True))
    return 1 if any(row.get("status") == "failed" for row in rows) else 0


if __name__ == "__main__":
    raise SystemExit(main())
