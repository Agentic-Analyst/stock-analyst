"""TTM data must update the model without corrupting fiscal-year history."""

import logging

import openpyxl
import pytest

from src.agents.fm.assumption_grounding import (
    _rolling_consensus_revenue_path,
    ground_assumptions,
)
from src.agents.fm.formula_evaluator import FormulaEvaluator
from src.agents.fm.tabs.tab_historical import HistoricalTabBuilder
from src.agents.fm.tabs.tab_assumptions import AssumptionsTabBuilder
from src.agents.fm.tabs.tab_projections import ProjectionsTabBuilder
from src.agents.fm.tabs.tab_raw import RawTabBuilder
from src.agents.fm.tabs.tab_sensitivity import SensitivityTabBuilder
from src.agents.fm.tabs.tab_summary import SummaryTabBuilder
from src.agents.fm.tabs.tab_valuation_exit_multiple_dcf import (
    ValuationExitMultipleDCFBuilder,
)
from src.agents.fm.tabs.tab_valuation_perpetual_growth_dcf import (
    ValuationPerpetualGrowthDCFBuilder,
)


def _current_bridge():
    return {
        "status": "current",
        "latest_period": "2026-06-30",
        "quarter_periods": ["2026-06-30", "2026-03-31", "2025-12-31", "2025-09-30"],
        "income_statement": {
            "Total Revenue": 400.0,
            "Gross Profit": 180.0,
            "Operating Income": 100.0,
        },
        "cash_flow": {"Depreciation And Amortization": 20.0},
        "normalized": {
            "capex_to_revenue": -0.08,
            "da_to_revenue": 0.05,
            "cash": 61.0,
            "short_term_investments": 29.0,
            "total_debt": 25.0,
        },
    }


def _grounding_data():
    return {
        "company_data": {
            "basic_info": {"currency": "USD", "country": "United States"},
            "market_data": {"market_cap": 1_000.0, "shares_outstanding": 10.0},
            "capital_structure": {"beta": 1.0, "total_debt": 25.0},
            "growth_profitability": {
                "operating_margins": 0.10,
                "ebitda_margins": 0.15,
                "gross_margins": 0.30,
            },
            "valuation_metrics": {},
        },
        "financial_statements": {
            "income_statement": {
                "2025-09-30": {
                    "Total Revenue": 360.0,
                    "Gross Profit": 144.0,
                    "Operating Income": 72.0,
                    "EBITDA": 90.0,
                },
                "2024-09-30": {
                    "Total Revenue": 340.0,
                    "Gross Profit": 136.0,
                    "Operating Income": 68.0,
                    "EBITDA": 85.0,
                },
            },
            "balance_sheet": {},
            "cash_flow": {},
        },
        "ttm_bridge": _current_bridge(),
    }


def test_grounding_prefers_ttm_margins_and_exposes_current_modeling_basis():
    grounded, notes = ground_assumptions(
        {
            "wacc": 0.09,
            "terminal_growth_rate": 0.025,
            "operating_margins": [0.10] * 5,
            "ebitda_margins": [0.15] * 5,
            "gross_margins": [0.30] * 5,
        },
        _grounding_data(),
    )

    assert grounded["operating_margins"][0] == pytest.approx(0.25)
    assert grounded["ebitda_margins"][0] == pytest.approx(0.30)
    assert grounded["gross_margins"][0] == pytest.approx(0.45)
    assert grounded["modeling_basis"]["period_end"] == "2026-06-30"
    assert grounded["modeling_basis"]["capex_to_revenue"] == pytest.approx(-0.08)
    assert grounded["modeling_basis"]["short_term_investments"] == pytest.approx(29.0)
    assert any("TTM/current balance-sheet bridge" in note for note in notes)


def test_projection_and_equity_bridge_use_current_ttm_inputs():
    basis = {**_current_bridge()["normalized"], "period_end": "2026-06-30"}
    workbook = openpyxl.Workbook()
    projections = ProjectionsTabBuilder(modeling_basis=basis).create_tab(workbook)
    valuation = ValuationPerpetualGrowthDCFBuilder(modeling_basis=basis).create_tab(workbook)

    assert "*0.050000000000" in projections["B12"].value
    assert "(-0.080000000000*1.00" in projections["B13"].value
    assert valuation["B30"].value == 61.0
    assert valuation["B31"].value == 25.0
    assert valuation["B32"].value == 29.0
    assert "2026-06-30" in valuation["G30"].value


def test_rolling_forecast_uses_ttm_revenue_and_current_period_endpoint():
    basis = {
        "revenue": 400.0,
        "forecast_basis": {
            "basis": "rolling_twelve_months",
            "period_end": "2026-06-30",
        },
    }
    workbook = openpyxl.Workbook()
    projections = ProjectionsTabBuilder(modeling_basis=basis).create_tab(workbook)

    assert projections["A1"].value == "PROJECTIONS (NTM1-NTM5)"
    assert projections["B2"].value == 2027
    assert projections["B3"].value == "=400.000000000000*(1+Assumptions!C7)"

    assumptions_workbook = openpyxl.Workbook()
    assumptions = AssumptionsTabBuilder({
        "forecast_basis": {
            **basis["forecast_basis"],
            "base_revenue": 400.0,
            "fiscal_year_progress": 0.75,
            "method": "fiscal-progress blend",
        },
    })
    assumptions.create_tab(assumptions_workbook)
    inputs = assumptions_workbook["Model_Inputs"]
    sheet = assumptions_workbook["Assumptions"]
    assert inputs["B1"].value == "NTM1"
    assert inputs["B12"].value == "Rolling twelve months"
    assert inputs["B14"].value == 400.0
    assert inputs["B15"].value == pytest.approx(0.75)
    assert sheet["B1"].value == "Latest FY Actual"
    assert sheet["C1"].value == "NTM1"

    valuation_workbook = openpyxl.Workbook()
    perpetual = ValuationPerpetualGrowthDCFBuilder(
        modeling_basis=basis
    ).create_tab(valuation_workbook)
    exit_tab = ValuationExitMultipleDCFBuilder(
        modeling_basis=basis
    ).create_tab(valuation_workbook)
    sensitivity = SensitivityTabBuilder(
        modeling_basis=basis
    ).create_tab(valuation_workbook)
    summary = SummaryTabBuilder(
        modeling_basis=basis
    ).create_tab(valuation_workbook)
    assert "NTM1-NTM10" in perpetual["A14"].value
    assert perpetual["A40"].value.startswith("Present EV / NTM5")
    assert exit_tab["A12"].value == "Terminal Year EBITDA (NTM10)"
    assert "NTM10 EBITDA" in sensitivity["A30"].value
    assert summary["A33"].value == "Revenue (NTM10)"
    assert summary["A52"].value == "PV of Explicit FCF (NTM1-NTM10)"


def test_rolling_consensus_path_blends_fiscal_endpoints_without_double_counting():
    data = _grounding_data()
    data["ttm_bridge"]["normalized"]["revenue"] = 400.0
    path = _rolling_consensus_revenue_path(
        data,
        {"0y": (440.0, 12), "+1y": (484.0, 10)},
        terminal_growth=0.025,
    )

    assert path is not None
    # 2026-06-30 is 273/365 of the way from the September 2025 fiscal close.
    progress = 273 / 365
    assert path["fiscal_year_progress"] == pytest.approx(progress)
    assert path["forecast_revenue"][0] == pytest.approx(
        440.0 * (1.0 - progress) + 484.0 * progress
    )
    assert path["revenue_growth_rates"][0] == pytest.approx(
        path["forecast_revenue"][0] / 400.0 - 1.0
    )
    assert path["analyst_counts"] == {"0y": 12, "+1y": 10}


def test_rolling_consensus_path_handles_leap_day_and_caps_close_grace():
    data = _grounding_data()
    data["ttm_bridge"]["normalized"]["revenue"] = 400.0
    data["financial_statements"]["income_statement"] = {
        "2024-02-29": data["financial_statements"]["income_statement"]["2025-09-30"],
    }
    data["ttm_bridge"]["latest_period"] = "2025-03-01"

    path = _rolling_consensus_revenue_path(
        data,
        {"0y": (440.0, 12), "+1y": (484.0, 10)},
        terminal_growth=0.025,
    )

    assert path is not None
    assert path["next_fiscal_end"] == "2025-02-28"
    assert path["fiscal_year_progress"] == 1.0
    assert path["forecast_revenue"][0] == pytest.approx(484.0)


@pytest.mark.parametrize(
    ("latest_period", "qualified"),
    [
        ("2025-09-22", {"0y": (440.0, 12), "+1y": (484.0, 10)}),
        ("2026-10-08", {"0y": (440.0, 12), "+1y": (484.0, 10)}),
        ("2026-06-30", {"0y": (440.0, 12)}),
    ],
)
def test_rolling_consensus_path_fails_closed_for_bad_clock_or_missing_anchor(
    latest_period, qualified,
):
    data = _grounding_data()
    data["ttm_bridge"]["normalized"]["revenue"] = 400.0
    data["ttm_bridge"]["latest_period"] = latest_period

    assert _rolling_consensus_revenue_path(
        data, qualified, terminal_growth=0.025,
    ) is None


def test_current_working_capital_base_requires_matching_field_definitions():
    data = _grounding_data()
    data["financial_statements"]["balance_sheet"] = {
        "2025-09-30": {
            "Accounts Receivable": 20.0,
            "Inventory": 10.0,
            "Accounts Payable": 15.0,
        },
        "2024-09-30": {
            "Accounts Receivable": 18.0,
            "Inventory": 9.0,
            "Accounts Payable": 14.0,
        },
    }
    data["financial_statements"]["income_statement"] = {
        "2025-09-30": {
            "Total Revenue": 365.0, "Cost Of Revenue": 182.5,
            "Operating Income": 72.0, "Gross Profit": 144.0,
        },
        "2024-09-30": {
            "Total Revenue": 330.0, "Cost Of Revenue": 165.0,
            "Operating Income": 68.0, "Gross Profit": 136.0,
        },
    }
    data["ttm_bridge"]["balance_sheet"] = {
        "Accounts Receivable": 24.0,
        "Inventory": 12.0,
        "Accounts Payable": 17.0,
    }

    grounded, notes = ground_assumptions({
        "wacc": 0.09, "terminal_growth_rate": 0.025,
        "dso_days": [30.0] * 5, "dio_days": [30.0] * 5,
        "dpo_days": [30.0] * 5,
    }, data)

    working_capital = grounded["modeling_basis"]["working_capital"]
    assert working_capital["net_working_capital"] == pytest.approx(19.0)
    assert working_capital["fields"] == {
        "accounts_receivable": "Accounts Receivable",
        "inventory": "Inventory",
        "accounts_payable": "Accounts Payable",
    }
    assert any("Current working-capital base" in note for note in notes)

    workbook = openpyxl.Workbook()
    projections = ProjectionsTabBuilder(
        modeling_basis=grounded["modeling_basis"]
    ).create_tab(workbook)
    assert projections["B18"].value == "=B17-(19.000000000000)"


def test_equity_bridge_prefers_verified_total_nonoperating_investments():
    basis = {
        **_current_bridge()["normalized"],
        "period_end": "2026-06-30",
        "non_operating_investments": 84.0,
        "non_operating_investments_source": "investments_and_advances",
    }
    workbook = openpyxl.Workbook()
    valuation = ValuationPerpetualGrowthDCFBuilder(
        modeling_basis=basis
    ).create_tab(workbook)

    assert valuation["B32"].value == 84.0
    assert "investments_and_advances" in valuation["G32"].value


def test_annual_fallback_adds_current_and_noncurrent_investment_buckets():
    period = "2025-09-30"
    statements = {
        "income_statement": {period: {
            "Total Revenue": 100.0,
            "Cost Of Revenue": 40.0,
            "Operating Income": 30.0,
            "Net Income": 25.0,
        }},
        "cash_flow": {period: {
            "Operating Cash Flow": 20.0,
            "Capital Expenditure": -5.0,
        }},
        "balance_sheet": {period: {
            "Cash And Cash Equivalents": 50.0,
            "Other Short Term Investments": 23.0,
            "Investments And Advances": 84.0,
            # Provider aliases repeat the broad bucket and must not themselves
            # be summed a second or third time.
            "Investmentin Financial Assets": 84.0,
            "Available For Sale Securities": 84.0,
            "Total Debt": 25.0,
        }},
    }
    workbook = openpyxl.Workbook()
    raw = RawTabBuilder()
    raw.add_data_from_json({"financial_statements": statements})
    raw.create_tab(workbook)
    HistoricalTabBuilder().create_tab(workbook)
    evaluator = FormulaEvaluator(workbook)
    evaluator.set_logger(logging.getLogger("test_investment_buckets"))
    cells = evaluator.evaluate_all_tabs()["Historical"]["cells"]

    assert cells["(31, 6)"] == 107.0
    assert cells["(37, 6)"] == -132.0


def test_combined_cash_and_investments_field_does_not_double_count_cash():
    period = "2025-09-30"
    statements = {
        "income_statement": {period: {
            "Total Revenue": 100.0,
            "Cost Of Revenue": 40.0,
            "Operating Income": 30.0,
            "Net Income": 25.0,
        }},
        "cash_flow": {period: {
            "Operating Cash Flow": 35.0,
            "Capital Expenditure": -10.0,
        }},
        "balance_sheet": {period: {
            "Cash And Cash Equivalents": 50.0,
            "Cash Cash Equivalents And Short Term Investments": 80.0,
            "Total Debt": 20.0,
            "Receivables": 40.0,
            "Payables And Accrued Expenses": 35.0,
            "Net PPE": 120.0,
        }},
    }
    workbook = openpyxl.Workbook()
    raw = RawTabBuilder()
    raw.add_data_from_json({"financial_statements": statements})
    raw.create_tab(workbook)
    HistoricalTabBuilder().create_tab(workbook)
    evaluator = FormulaEvaluator(workbook)
    evaluator.set_logger(logging.getLogger("test_ttm_modeling_bridge"))
    cells = evaluator.evaluate_all_tabs()["Historical"]["cells"]

    assert cells["(30, 6)"] == 50.0
    assert cells["(31, 6)"] == 30.0
    assert cells["(32, 6)"] == 40.0
    assert cells["(34, 6)"] == 35.0
    assert cells["(37, 6)"] == -60.0
    assert cells["(38, 6)"] == 120.0


def test_missing_latest_sga_breakout_uses_reported_operating_income_residual():
    period = "2026-01-31"
    statements = {
        "income_statement": {period: {
            "Total Revenue": 713.0,
            "Cost Of Revenue": 535.0,
            "Gross Profit": 178.0,
            "Research And Development": None,
            "Selling General And Administration": None,
            "Other Operating Expenses": 148.0,
            "Operating Income": 30.0,
            "Net Income": 22.0,
        }},
        "cash_flow": {period: {
            "Operating Cash Flow": 42.0,
            "Capital Expenditure": -27.0,
        }},
        "balance_sheet": {period: {"Cash And Cash Equivalents": 11.0}},
    }
    workbook = openpyxl.Workbook()
    raw = RawTabBuilder()
    raw.add_data_from_json({"financial_statements": statements})
    raw.create_tab(workbook)
    HistoricalTabBuilder().create_tab(workbook)
    evaluator = FormulaEvaluator(workbook)
    evaluator.set_logger(logging.getLogger("test_missing_sga_breakout"))
    cells = evaluator.evaluate_all_tabs()["Historical"]["cells"]

    assert cells["(7, 6)"] == pytest.approx(148.0)
    assert cells["(44, 6)"] == pytest.approx(30.0)
    assert cells["(45, 6)"] == pytest.approx(0.0)


def test_working_capital_drivers_do_not_mix_quarterly_and_annual_field_scopes():
    data = _grounding_data()
    data["financial_statements"]["income_statement"] = {
        "2025-09-30": {"Total Revenue": 365.0, "Cost Of Revenue": 182.5},
        "2024-09-30": {"Total Revenue": 330.0, "Cost Of Revenue": 165.0},
    }
    data["financial_statements"]["balance_sheet"] = {
        "2025-09-30": {
            "Receivables": 36.5, "Inventory": 18.25,
            "Payables And Accrued Expenses": 73.0,
        },
        "2024-09-30": {
            "Receivables": 33.0, "Inventory": 16.5,
            "Payables And Accrued Expenses": 66.0,
        },
    }
    # A current quarter whose broad receivable and narrow payable fields are
    # not comparable with the annual definitions must not seed FY1.
    data["ttm_bridge"]["income_statement"].update({"Cost Of Revenue": 200.0})
    data["ttm_bridge"]["balance_sheet"] = {
        "Receivables": 300.0, "Payables": 5.0,
    }

    grounded, _ = ground_assumptions({
        "wacc": 0.09,
        "terminal_growth_rate": 0.025,
        "dso_days": [30.0] * 5,
        "dio_days": [30.0] * 5,
        "dpo_days": [30.0] * 5,
    }, data)

    assert grounded["dso_days"][0] == pytest.approx(36.5 / 365.0 * 365.0)
    assert grounded["dio_days"][0] == pytest.approx(18.25 / 182.5 * 365.0)
    assert grounded["dpo_days"][0] == pytest.approx(73.0 / 182.5 * 365.0)
    assert "working_capital" not in grounded["modeling_basis"]

    workbook = openpyxl.Workbook()
    projections = ProjectionsTabBuilder(
        modeling_basis=grounded["modeling_basis"]
    ).create_tab(workbook)
    assert projections["B18"].value == (
        "=B17-(Historical!$F$32+Historical!$F$33-Historical!$F$34)"
    )


def test_working_capital_prefers_narrow_trade_balances_when_broad_totals_exist():
    data = _grounding_data()
    data["financial_statements"]["income_statement"] = {
        "2025-12-31": {
            "Total Revenue": 365.0, "Cost Of Revenue": 182.5,
            "Operating Income": 90.0, "Net Income": 70.0,
        },
        "2024-12-31": {
            "Total Revenue": 330.0, "Cost Of Revenue": 165.0,
            "Operating Income": 80.0, "Net Income": 60.0,
        },
    }
    data["financial_statements"]["cash_flow"] = {
        "2025-12-31": {"Operating Cash Flow": 85.0},
        "2024-12-31": {"Operating Cash Flow": 75.0},
    }
    data["financial_statements"]["balance_sheet"] = {
        "2025-12-31": {
            "Accounts Receivable": 20.0, "Receivables": 30.0,
            "Inventory": 10.0,
            "Accounts Payable": 15.0, "Payables": 18.0,
            "Payables And Accrued Expenses": 80.0,
            "Cash And Cash Equivalents": 40.0,
        },
        "2024-12-31": {
            "Accounts Receivable": 18.0, "Receivables": 28.0,
            "Inventory": 9.0,
            "Accounts Payable": 14.0, "Payables": 17.0,
            "Payables And Accrued Expenses": 75.0,
            "Cash And Cash Equivalents": 35.0,
        },
    }

    grounded, _ = ground_assumptions({
        "wacc": 0.09, "terminal_growth_rate": 0.025,
        "dso_days": [30.0] * 5, "dio_days": [30.0] * 5,
        "dpo_days": [30.0] * 5,
    }, data)

    assert grounded["dso_days"][0] == pytest.approx(20.0)
    assert grounded["dpo_days"][0] == pytest.approx(30.0)

    workbook = openpyxl.Workbook()
    raw = RawTabBuilder()
    raw.add_data_from_json({"financial_statements": data["financial_statements"]})
    raw.create_tab(workbook)
    HistoricalTabBuilder().create_tab(workbook)
    evaluator = FormulaEvaluator(workbook)
    evaluator.set_logger(logging.getLogger("test_working_capital_narrow_lines"))
    cells = evaluator.evaluate_all_tabs()["Historical"]["cells"]
    assert cells["(32, 6)"] == 20.0
    assert cells["(34, 6)"] == 15.0
