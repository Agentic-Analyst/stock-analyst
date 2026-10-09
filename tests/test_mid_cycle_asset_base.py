"""
A memory maker's D&A and capex follow its asset base, not the price cycle.

The engine scales both with projected revenue. For a memory maker in a boom
that revenue is price: Micron's trailing capex (28% of revenue that had
tripled) became $86B and $84B a year in the Street's covered years, more than
three times what it was spending, building plant no mid-cycle year earns on,
while D&A on the trend years stayed at the boom's 10% of revenue. Now both are
measured against trend revenue and grow with it, capex converging to the
engine's usual 1.1x D&A by FY5; and the exit multiple is held to what the
model's own terminal cash conversion justifies, because today's EV/EBITDA
prices the same cycle.

Run:  python -m pytest tests/test_mid_cycle_asset_base.py -q
"""

import os
import sys

import openpyxl
import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "src"))
sys.path.insert(0, os.path.join(_ROOT, "tests"))

from src.agents.fm import memory_cycle  # noqa: E402
from src.agents.fm.tabs.tab_projections import ProjectionsTabBuilder  # noqa: E402
from src.agents.fm.tabs.tab_valuation_exit_multiple_dcf import (  # noqa: E402
    ValuationExitMultipleDCFBuilder,
)
from src.reinvestment_sensitivity import build_reinvestment_sensitivity  # noqa: E402
from test_memory_mid_cycle import _grounded, _sandisk  # noqa: E402


def _revenue(grounded):
    revenue = [grounded["modeling_basis"]["revenue"]]
    for rate in grounded["revenue_growth_rates"]:
        revenue.append(revenue[-1] * (1 + rate))
    return revenue


def test_da_and_capex_follow_trend_revenue_not_the_price(monkeypatch):
    monkeypatch.setenv("RISK_FREE_USD", "0.04")
    grounded, notes = _grounded(_sandisk())
    basis = grounded["modeling_basis"]
    inputs = memory_cycle.mid_cycle_inputs(_sandisk())
    period = grounded["mid_cycle"]["base_period_end"]
    revenue = _revenue(grounded)

    # The trailing dollars, measured against the trend revenue at the base
    # period instead of the boom revenue they were reported on.
    trend_now = memory_cycle.trend_revenue(inputs, 0, period)
    da_share = basis["da_to_revenue"] * basis["revenue"] / trend_now
    capex_share = basis["capex_to_revenue"] * basis["revenue"] / trend_now
    asset = grounded["mid_cycle"]["asset_base"]
    assert asset["da_share_of_trend_revenue"] == pytest.approx(da_share)
    assert asset["capex_share_of_trend_revenue"] == pytest.approx(capex_share)

    for i in range(5):
        trend = memory_cycle.trend_revenue(inputs, i + 1, period)
        da = basis["da_to_revenue_by_year"][i] * revenue[i + 1]
        capex = basis["capex_to_revenue_by_year"][i] * revenue[i + 1]
        assert da == pytest.approx(da_share * trend)
        # The engine's glide, on trend dollars: from today's intensity to
        # 1.1x D&A by FY5.
        assert capex == pytest.approx(trend * (capex_share * (1 - i / 4) - 1.1 * da_share * i / 4))
        # EBITDA is EBIT plus the D&A actually projected that year.
        assert grounded["ebitda_margins"][i] == pytest.approx(
            grounded["operating_margins"][i] + basis["da_to_revenue_by_year"][i])
    assert capex == pytest.approx(-1.1 * da)
    assert any("D&A and capex on the asset base" in note for note in notes)


def test_without_trailing_ratios_the_engine_path_is_kept(monkeypatch):
    monkeypatch.setenv("RISK_FREE_USD", "0.04")
    payload = _sandisk()
    payload["ttm_bridge"]["normalized"].pop("da_to_revenue", None)
    grounded, notes = _grounded(payload)
    basis = grounded.get("modeling_basis") or {}
    if "da_to_revenue" in basis:
        pytest.skip("fixture supplies D&A from elsewhere")
    assert "da_to_revenue_by_year" not in basis
    assert grounded["mid_cycle"]["asset_base"] is None
    assert any("no trailing ratios" in note for note in notes)


@pytest.mark.parametrize("da,capex", [(None, -0.1), (0.1, None), (0.7, -0.1), (0.1, 0.05),
                                      (float("nan"), -0.1), (True, -0.1)])
def test_implausible_trailing_ratios_leave_the_engine_path(da, capex):
    inputs = memory_cycle.mid_cycle_inputs(_sandisk())
    basis = {"da_to_revenue": da, "capex_to_revenue": capex}
    assert memory_cycle.asset_base_path(
        basis, inputs, inputs["base_period_end"], 20e9, [20e9] * 5) is None


def test_a_share_of_trend_outside_any_asset_base_is_a_data_problem():
    inputs = memory_cycle.mid_cycle_inputs(_sandisk())
    trend_now = memory_cycle.trend_revenue(inputs, 0, inputs["base_period_end"])
    # D&A reported on revenue three times trend would be 90% of trend revenue.
    basis = {"da_to_revenue": 0.3, "capex_to_revenue": -0.1}
    assert memory_cycle.asset_base_path(
        basis, inputs, inputs["base_period_end"], 3 * trend_now, [3 * trend_now] * 5) is None


def _projections(basis):
    workbook = openpyxl.Workbook()
    ws = ProjectionsTabBuilder(modeling_basis=basis).create_tab(workbook)
    return [ws.cell(row=12, column=c).value for c in range(2, 7)], \
        [ws.cell(row=13, column=c).value for c in range(2, 7)]


def test_the_workbook_uses_one_ratio_per_year_when_both_are_given():
    by_year = {
        "da_to_revenue": 0.10, "capex_to_revenue": -0.28,
        "da_to_revenue_by_year": [0.04, 0.035, 0.1, 0.28, 0.28],
        "capex_to_revenue_by_year": [-0.12, -0.08, -0.19, -0.42, -0.31],
    }
    da, capex = _projections(by_year)
    assert da[0] == "=B3*0.040000000000" and da[3] == "=E3*0.280000000000"
    assert capex[0] == "=B3*-0.120000000000" and capex[4] == "=F3*-0.310000000000"

    # Both or neither: one list alone, a short list or an implausible value
    # keeps the engine's single-ratio path for both rows.
    for broken in ({"capex_to_revenue_by_year": None},
                   {"da_to_revenue_by_year": [0.04] * 4},
                   {"capex_to_revenue_by_year": [-0.12, 0.3, -0.19, -0.42, -0.31]}):
        da, capex = _projections({**by_year, **broken})
        assert da[0] == "=B3*0.100000000000"
        assert capex[0].startswith("=B3*(-0.280000000000*1.00")


def test_a_mid_cycle_exit_multiple_is_held_to_its_ceiling_whatever_the_conversion():
    def multiple_formula(always_cap):
        workbook = openpyxl.Workbook()
        ws = ValuationExitMultipleDCFBuilder(
            exit_multiple=10.4, growth_cap=0.04, always_cap=always_cap,
        ).create_tab(workbook)
        return ws.cell(row=13, column=2).value

    gated = multiple_formula(False)
    assert "$K$7/$B$12>=0.30" in gated
    held = multiple_formula(True)
    assert "$K$7/$B$12>=0.30" not in held and "$K$7>0" in held
    assert "MIN($B$3,($K$7/$B$12)" in held


def test_the_run_rate_capex_case_does_not_apply_to_an_asset_base_path():
    result = build_reinvestment_sensitivity(
        {}, {}, modeling_basis={"capex_to_revenue": -0.28, "da_to_revenue": 0.1,
                                "capex_to_revenue_by_year": [-0.12] * 5})
    assert result["status"] == "unavailable"
    assert "asset base" in result["reason"]


# ── Micron's own workbook, built offline from its 9 October 2026 data ───────

def _build_micron():
    import pathlib
    from src.agents.fm.financial_model_builder import FinancialModelBuilder
    from src.agents.fm.formula_evaluator import model_integrity

    class Quiet:
        def __getattr__(self, _name):
            return lambda *args, **kwargs: None

    builder = FinancialModelBuilder("MU", Quiet())
    builder.load_json_file(pathlib.Path(__file__).parent / "fixtures"
                           / "mu_memory_boom_2026_10_09.json")
    builder.build_model()
    results = builder.formula_evaluator.evaluate_all_tabs()
    return builder, results, model_integrity(results)


def test_micron_builds_plant_for_its_trend_not_its_boom():
    builder, results, integrity = _build_micron()
    assert integrity["status"] == "ready", integrity.get("issues")
    basis = builder.llm_assumptions["modeling_basis"]
    projections = results["Projections"]["cells"]
    ttm_capex = -basis["capex_to_revenue"] * basis["revenue"]
    capex = [-projections[f"(13, {col})"] for col in range(2, 7)]
    # About what Micron spent over the last twelve months, never three times it.
    assert ttm_capex < capex[0] < 1.3 * ttm_capex
    assert capex[1] < 1.3 * ttm_capex
    # The boom path took net PP&E from $47B to $167B (3.5x) on capex scaled
    # with price; on the asset base it ends near $101B, the trailing build-out
    # gliding to maturity.
    assert projections["(46, 6)"] < 2.5 * projections["(43, 2)"]
    # D&A on the trend's asset base: about a quarter of FY5 revenue, not the
    # boom's 10%.
    assert projections["(12, 6)"] / projections["(3, 6)"] > 0.2
    assert projections["(13, 6)"] == pytest.approx(-1.1 * projections["(12, 6)"])


def test_micron_exit_leg_is_held_to_its_own_economics():
    _builder, results, _integrity = _build_micron()
    exit_tab = results["Valuation (Exit Multiple)"]["cells"]
    multiple, reference = exit_tab["(13, 2)"], exit_tab["(3, 2)"]
    # Its terminal conversion is under the 30% gate, legitimately; held to the
    # ceiling anyway, not today's multiple on boom EBITDA.
    assert multiple < reference
    perpetual = results["Valuation (DCF)"]["cells"]["(37, 2)"]
    exit_value = exit_tab["(25, 2)"] if isinstance(exit_tab.get("(25, 2)"), (int, float)) else None
    if exit_value is not None:
        assert abs(exit_value / perpetual - 1) < 0.05
