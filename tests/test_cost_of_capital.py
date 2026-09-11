"""
The discount rate must be derived, and the derivation must be visible.

THE BUG. A user compared our LVMH report against one a frontier model wrote.
Ours said SELL at EUR 390; theirs said HOLD at EUR 480 against a EUR 452.50
market price. The entire disagreement was the discount rate.

src/agents/fm/tabs/tab_assumptions.py wrote literal constants into the cells the
DCF reads — Rf 4.5%, ERP 6.5%, levered beta 1.2, Kd 5.5%, equity weight 85% —
for every company on earth. LVMH's observed beta is 0.84 and it is a euro
issuer, so its risk-free rate is the Bund, not a US Treasury. Those constants
produced a WACC of 11.01% and a fair value of EUR 269 on projections that return
roughly EUR 433 at 8%.

Two aggravating facts. compute_capm_wacc already derived a defensible 8.57% and
threw the components away, so nothing could be written into the cells. And the
report PRINTED that derived rate while the workbook discounted at 11.01% — the
reader was shown a number the model never used.

Run:  python -m pytest tests/test_cost_of_capital.py -q
"""

import os
import sys

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "src"))

from src.agents.fm.assumption_grounding import (
    capm_components,
    compute_capm_wacc,
    risk_free_details,
    risk_free_rate,
)
from src.agents.fm import assumption_grounding as _ag
from src.agents.fm.sovereign_rates import _SNAPSHOT

# LVMH as scraped, the run that exposed this.
LVMH = {
    "basic_info": {"currency": "EUR"},
    "capital_structure": {"beta": 0.842, "total_debt": 36_731_000_000},
    "market_data": {"market_cap": 226_443_575_296},
    "growth_profitability": {},
}
US_MEGACAP = {
    "basic_info": {"currency": "USD"},
    "capital_structure": {"beta": 1.05, "total_debt": 30_000_000_000},
    "market_data": {"market_cap": 3_000_000_000_000},
    "growth_profitability": {},
}


class TestRiskFreeByCurrency:
    # The suite runs with the feeds switched off (conftest), so every rate here
    # is the dated snapshot — the path a run takes when FRED or the ECB is down.

    def test_euro_issuer_does_not_get_a_us_treasury(self):
        """
        Discounting euro cash flows at a US Treasury yield is a currency
        mismatch in the cost of capital — the same class of error as quoting a
        EUR company in dollars.
        """
        rate, source = risk_free_rate("EUR")
        assert rate == pytest.approx(_SNAPSHOT["EUR"][0])     # Germany is Aaa: nothing subtracted
        assert "euro-area" in source or "Bund" in source
        assert "US 10Y" not in source

    def test_dated_rates_disclose_their_date(self):
        _, source = risk_free_rate("EUR")
        assert "as of" in source

    def test_a_yen_issuer_gets_the_jgb_not_a_treasury(self):
        d = risk_free_details("JPY")
        assert d["sovereign_yield"] == pytest.approx(_SNAPSHOT["JPY"][0])
        assert "JGB" in d["label"]

    def test_unknown_currency_is_labelled_a_proxy_not_passed_off_as_local(self):
        rate, source = risk_free_rate("XXX")
        assert rate > 0
        assert "proxy" in source.lower()
        assert "XXX" in source

    def test_operator_can_override_without_a_deploy(self, monkeypatch):
        monkeypatch.setenv("RISK_FREE_EUR", "0.0290")
        rate, source = risk_free_rate("EUR")
        assert rate == pytest.approx(0.0290)
        assert "override" in source

    def test_a_nonsense_override_is_ignored(self, monkeypatch):
        monkeypatch.setenv("RISK_FREE_EUR", "not-a-number")
        rate, _ = risk_free_rate("EUR")
        assert rate == pytest.approx(_SNAPSHOT["EUR"][0])

    def test_missing_currency_defaults_to_usd(self):
        _, source = risk_free_rate(None)
        assert "US 10Y" in source


INDIAN_MIDCAP = {
    "basic_info": {"currency": "INR", "country": "India"},
    "capital_structure": {"beta": 1.0, "total_debt": 1_000_000_000},
    "market_data": {"market_cap": 5_000_000_000},
    "growth_profitability": {},
}


class TestRiskFreeBuild:
    """
    Rf = government bond yield in the currency − the sovereign's default spread.

    A Baa3 bond is not default-free; its default risk is what the country
    premium prices into Ke. Keeping it in Rf as well charged India twice.
    """

    def test_a_rated_sovereign_s_default_spread_is_taken_out(self):
        d = risk_free_details("INR")
        assert d["sovereign_yield"] == pytest.approx(_SNAPSHOT["INR"][0])
        assert 0.015 < d["default_spread"] < 0.025          # India, Baa3
        assert d["rate"] == pytest.approx(d["sovereign_yield"] - d["default_spread"])
        assert "default spread" in d["label"] and "Baa3" in d["label"]
        assert f"= {d['rate']*100:.2f}%" in d["label"]

    def test_an_aaa_sovereign_is_used_as_is(self):
        d = risk_free_details("EUR")
        assert d["default_spread"] == 0.0
        assert d["rate"] == d["sovereign_yield"]
        assert "default spread" not in d["label"]

    def test_the_us_has_been_aa1_since_may_2025(self):
        d = risk_free_details("USD")
        assert 0.001 < d["default_spread"] < 0.005
        assert "Aa1" in d["label"]

    def test_a_proxy_is_still_a_us_bond(self):
        """It carries the US spread, not the target currency's; see TestReviewFindings."""
        d = risk_free_details("XXX")
        assert d["proxy"] is True
        assert d["default_spread"] == pytest.approx(risk_free_details("USD")["default_spread"])

    def test_cost_of_debt_is_priced_off_the_issuer_s_own_government(self):
        """An Indian issuer borrows above the G-Sec, not above the G-Sec less India's spread."""
        c = capm_components(INDIAN_MIDCAP)
        assert c["pre_tax_cost_of_debt"] == pytest.approx(c["sovereign_yield"] + 0.015)   # = G-Sec + spread
        assert c["pre_tax_cost_of_debt"] == pytest.approx(c["risk_free_rate"] + c["domicile_default_spread"] + 0.015)
        assert c["domicile"] == "India"
        assert c["kd_source"].startswith("risk-free 5.02% (default-free) + India sovereign spread 1.87%")
        assert c["kd_source"].endswith("+ 1.5% credit spread")

    def test_the_country_premium_is_the_published_one(self):
        c = capm_components(INDIAN_MIDCAP)
        assert 0.02 <= c["country_risk_premium"] <= 0.04
        assert "Damodaran" in c["crp_source"]
        assert c["equity_risk_premium_total"] == pytest.approx(c["equity_risk_premium"] + c["country_risk_premium"])
        assert c["cost_of_equity"] == pytest.approx(c["risk_free_rate"] + c["beta"] * c["equity_risk_premium_total"])

    def test_a_us_issuer_s_cost_of_equity_barely_moves_at_beta_one(self):
        """
        Subtracting the US default spread from Rf and adding the same spread's
        premium to the ERP cancel at beta 1: the methodology is consistent,
        not a hidden repricing of every US stock.
        """
        c = capm_components(US_MEGACAP)
        naive = c["sovereign_yield"] + c["beta"] * c["equity_risk_premium"]
        assert abs(c["cost_of_equity"] - naive) < 0.001

    def test_mature_erp_can_be_overridden(self, monkeypatch):
        monkeypatch.setenv("EQUITY_RISK_PREMIUM", "0.05")
        assert capm_components(US_MEGACAP)["equity_risk_premium"] == pytest.approx(0.05)
        # An out-of-band override is ignored. The fallback used to be the 5.5%
        # house constant; it is now Damodaran's published implied premium,
        # which is the point of the change — so assert the fallback IS the
        # engine's resolved figure rather than re-pinning a literal.
        monkeypatch.setenv("EQUITY_RISK_PREMIUM", "0.5")
        monkeypatch.delenv("EQUITY_RISK_PREMIUM", raising=False)
        _fallback = _ag._mature_erp()
        monkeypatch.setenv("EQUITY_RISK_PREMIUM", "0.5")
        assert capm_components(US_MEGACAP)["equity_risk_premium"] == pytest.approx(_fallback)

    def test_the_build_is_published_for_the_workbook_and_report(self):
        c = capm_components(INDIAN_MIDCAP)
        for key in ("sovereign_yield", "sovereign_default_spread", "risk_free_as_of",
                    "risk_free_kind", "kd_source", "crp_source"):
            assert key in c, key
        assert c["risk_free_kind"] == "snapshot"


class TestCapmBuild:
    def test_it_returns_the_derivation_not_just_the_answer(self):
        """
        The components existed as locals and were discarded, which is why the
        workbook had to invent its own constants.
        """
        c = capm_components(LVMH)
        for key in ("risk_free_rate", "risk_free_source", "beta", "beta_source",
                    "equity_risk_premium", "cost_of_equity", "pre_tax_cost_of_debt",
                    "after_tax_cost_of_debt", "equity_weight", "debt_weight", "wacc"):
            assert key in c, key

    def test_beta_is_the_company_s_own_not_1_point_2(self):
        c = capm_components(LVMH)
        # Blume: 0.67 * 0.842 + 0.33
        assert c["beta"] == pytest.approx(0.67 * 0.842 + 0.33, abs=1e-6)
        assert "0.84" in c["beta_source"]

    def test_a_low_beta_euro_issuer_is_not_discounted_at_eleven_percent(self):
        """The shipped model returned 11.01% here and a EUR 269 fair value."""
        c = capm_components(LVMH)
        assert 0.065 <= c["wacc"] <= 0.090, c["wacc"]

    def test_weights_come_from_the_actual_capital_structure(self):
        c = capm_components(LVMH)
        expected_w_d = 36_731_000_000 / (226_443_575_296 + 36_731_000_000)
        assert c["debt_weight"] == pytest.approx(expected_w_d, abs=1e-6)
        assert c["equity_weight"] + c["debt_weight"] == pytest.approx(1.0)

    def test_a_us_company_still_gets_a_us_risk_free(self):
        c = capm_components(US_MEGACAP)
        assert c["currency"] == "USD"
        assert "US 10Y" in c["risk_free_source"]

    def test_missing_beta_falls_back_to_the_market(self):
        c = capm_components({"basic_info": {"currency": "USD"},
                             "capital_structure": {}, "market_data": {}})
        assert c["beta"] == 1.0
        assert "no observed beta" in c["beta_source"]

    def test_the_note_still_works_for_existing_callers(self):
        wacc, note = compute_capm_wacc(LVMH)
        assert 0.0 < wacc < 0.2
        assert "CAPM WACC" in note

    def test_clamping_is_disclosed(self):
        """A clamped WACC is a modelling decision, not a derivation."""
        wild = {"basic_info": {"currency": "USD"},
                "capital_structure": {"beta": 5.0, "total_debt": 0},
                "market_data": {"market_cap": 1_000_000_000}}
        c = capm_components(wild)
        if c["wacc_clamped"]:
            _, note = compute_capm_wacc(wild)
            assert "clamped" in note


class TestWorkbookUsesTheDerivation:
    def test_the_assumptions_tab_no_longer_hardcodes_capm_inputs(self):
        """The literals that produced 11.01% for every company on earth."""
        source = open(os.path.join(_ROOT, "src", "agents", "fm", "tabs",
                                   "tab_assumptions.py"), encoding="utf-8").read()
        dcf_block = source[source.index("_setup_dcf_parameters"):]
        for literal in ("value=0.045", "value=0.065", "value=1.2", "value=0.055",
                        "value=0.85"):
            assert literal not in dcf_block, literal

    def test_the_assumptions_tab_reads_the_capm_build(self):
        source = open(os.path.join(_ROOT, "src", "agents", "fm", "tabs",
                                   "tab_assumptions.py"), encoding="utf-8").read()
        assert 'self.llm_assumptions.get("capm")' in source

    def test_grounding_publishes_the_build_for_the_workbook(self):
        import inspect
        from src.agents.fm import assumption_grounding
        source = inspect.getsource(assumption_grounding.ground_assumptions)
        assert 'a["capm"] = capm' in source

    def test_seeded_cells_reproduce_the_derived_wacc(self):
        """
        The workbook computes Ke = Rf + beta*ERP and blends it with after-tax Kd
        at the capital-structure weights. Seeding those five cells must land on
        the same WACC the grounding computed, or the workbook and the report
        disagree again.
        """
        c = capm_components(LVMH)
        ke = c["risk_free_rate"] + c["beta"] * c["equity_risk_premium_total"]
        rebuilt = (c["equity_weight"] * ke
                   + c["debt_weight"] * c["after_tax_cost_of_debt"])
        assert rebuilt == pytest.approx(c["wacc"], abs=1e-6)


class TestLeasesAreNotChargedTwice:
    def test_fcf_does_not_deduct_lease_repayments(self):
        """
        A tempting fix that is wrong here. This model treats leases as debt: it
        adds back total D&A (including right-of-use depreciation) and its equity
        bridge deducts balance-sheet Total Debt, which INCLUDES capital lease
        obligations — EUR 16.4bn of LVMH's EUR 36.7bn. Deducting lease principal
        from the cash flow as well charges the same obligation twice.
        """
        source = open(os.path.join(_ROOT, "src", "agents", "fm", "tabs",
                                   "tab_projections.py"), encoding="utf-8").read()
        fcf = source[source.index("def _setup_fcf_row"):
                     source.index("def _setup_analytics_section")]
        assert "Assumptions!$B$32" not in fcf
        assert "leases as debt" in fcf

    def test_the_equity_bridge_still_deducts_total_debt(self):
        """The other half of the convention. If this changes, the FCF must too."""
        source = open(os.path.join(_ROOT, "src", "agents", "fm", "tabs",
                                   "tab_valuation_perpetual_growth_dcf.py"),
                      encoding="utf-8").read()
        assert "Total Debt" in source


class TestReportShowsTheDerivation:
    """
    The report printed a bare WACC and asserted it was reasonable. Worse, it
    printed LLM_Inferred's rate while the workbook discounted with the
    Assumptions cells — one number shown, another used.
    """

    def _computed(self, labelled=True):
        cells = {
            '(3, 2)': 0.0324, '(4, 2)': 0.055, '(5, 2)': 0.894,
            '(6, 2)': 0.0816, '(7, 2)': 0.0474, '(8, 2)': 0.3279,
            '(9, 2)': 0.0319, '(10, 2)': 0.8604, '(11, 2)': 0.1396,
            '(12, 2)': 0.0746,
        }
        if labelled:
            cells.update({
                '(3, 1)': 'Risk-Free Rate (Rf)', '(4, 1)': 'Equity Risk Premium (ERP)',
                '(5, 1)': 'Levered Beta (β)', '(6, 1)': 'Cost of Equity (Ke)',
                '(7, 1)': 'Pre-Tax Cost of Debt (Kd)', '(8, 1)': 'Tax Rate (T)',
                '(9, 1)': 'After-Tax Cost of Debt', '(10, 1)': 'Equity Weight (E/V)',
                '(11, 1)': 'Debt Weight (D/V)', '(12, 1)': 'WACC',
            })
        return {'Valuation (DCF)': {'cells': cells}}

    def test_a_shifted_row_does_not_silently_read_the_wrong_number(self):
        """Labels win over positions; a moved row must not become a wrong rate."""
        from src.report_agent import extract_cost_of_capital
        cells = self._computed()['Valuation (DCF)']['cells']
        shifted = {}
        for key, value in cells.items():
            row, col = eval(key)
            shifted[f'({row + 5}, {col})'] = value
        coc = extract_cost_of_capital({'Valuation (DCF)': {'cells': shifted}})
        assert coc['wacc'] == pytest.approx(0.0746)
        assert coc['beta'] == pytest.approx(0.894)

    def test_unlabelled_models_still_parse_by_row(self):
        from src.report_agent import extract_cost_of_capital
        coc = extract_cost_of_capital(self._computed(labelled=False))
        assert coc['wacc'] == pytest.approx(0.0746)

    def test_reads_the_rate_the_dcf_actually_used(self):
        from src.report_agent import extract_cost_of_capital
        coc = extract_cost_of_capital(self._computed())
        assert coc['wacc'] == pytest.approx(0.0746)
        assert coc['risk_free_rate'] == pytest.approx(0.0324)
        assert coc['beta'] == pytest.approx(0.894)

    def test_missing_tab_degrades_quietly(self):
        from src.report_agent import extract_cost_of_capital
        coc = extract_cost_of_capital({})
        assert coc['wacc'] is None

    def test_non_numeric_cells_are_dropped(self):
        from src.report_agent import extract_cost_of_capital
        coc = extract_cost_of_capital(
            {'Valuation (DCF)': {'cells': {'(12, 2)': 'A. WACC CALCULATION'}}})
        assert coc['wacc'] is None


class TestSensitivityGrid:
    PROJECTIONS = {'fcf': [15.34e9, 14.93e9, 14.35e9, 13.71e9, 12.96e9]}
    # The grid runs the DCF tab's own model: ten FCF periods and its bridge.
    VALUATION = {
        'dcf_perpetual': {'enterprise_value': 145.7e9, 'equity_value': 131.3e9,
                          'intrinsic_value_per_share': 263.7},
        'dcf_inputs': {
            'fcf': [15.34e9, 14.93e9, 14.35e9, 13.71e9, 12.96e9,
                    13.28e9, 13.62e9, 13.96e9, 14.31e9, 14.67e9],
            'cash': 8.794e9, 'debt': 36.731e9, 'investments': 13.502e9,
            'shares': 497_976_118,
        },
    }

    def _grid(self, wacc=0.0746, g=0.025):
        from src.report_agent import build_sensitivity_grid
        return build_sensitivity_grid(self.PROJECTIONS, g, self.VALUATION, wacc)

    def test_the_grid_is_recomputed_not_read_from_the_workbook(self):
        """
        The Sensitivity tab exists but its interior cells carry formatting and
        no values, so there is nothing to read back.
        """
        grid = self._grid()
        assert grid
        assert grid.count("\n") >= 6          # header, rule, five rows

    def test_the_base_case_row_is_marked(self):
        grid = self._grid(wacc=0.0746)
        assert "**7.46%**" in grid

    def test_value_falls_as_the_discount_rate_rises(self):
        grid = self._grid()
        rows = [r for r in grid.splitlines() if r.startswith("|") and "%" in r][1:]
        first = int(rows[0].split("|")[3].strip().replace(",", ""))
        last = int(rows[-1].split("|")[3].strip().replace(",", ""))
        assert last < first

    def test_a_diverging_terminal_value_is_marked_not_printed(self):
        """g >= WACC makes the Gordon formula meaningless, not merely large."""
        grid = self._grid(wacc=0.03, g=0.035)
        assert "n/m" in grid

    def test_missing_inputs_yield_no_grid_rather_than_a_wrong_one(self):
        from src.report_agent import build_sensitivity_grid
        no_inputs = {'dcf_perpetual': self.VALUATION['dcf_perpetual']}
        assert build_sensitivity_grid(self.PROJECTIONS, 0.025, no_inputs, 0.08) == ""
        assert build_sensitivity_grid(self.PROJECTIONS, 0.025, {}, 0.08) == ""

    def test_the_prompt_asks_which_input_dominates(self):
        """
        The point of the grid: on these projections the value moves far more per
        50bp of WACC than per point of growth, and the reader should be told.
        """
        prompt = open(os.path.join(_ROOT, "prompts", "report_valuation.md"),
                      encoding="utf-8").read()
        assert "{tables}" in prompt
        assert "Do NOT reproduce" in prompt
        assert "which input dominates" in prompt


class TestReportAssemblyRuns:
    """
    Two NameErrors shipped from this area in one day: `validation_report` in the
    recommendation formatter (eleven days live, hidden by a bare except), and a
    reference to `assumptions` in integrate_report_sections placed a hundred
    lines above where that local is bound.

    Source inspection catches neither reliably. Assembling a report does.
    """

    def _data(self):
        five = [0.03, 0.028, 0.026, 0.024, 0.022]
        return {
            'company_overview': {
                'company_name': 'LVMH', 'ticker': 'MC.PA', 'sector': 'Consumer Cyclical',
                'industry': 'Luxury Goods', 'current_price': 459.35,
                'market_cap': 226_443_575_296, 'currency': 'EUR',
                'exchange': 'PAR', 'country': 'France', 'website': 'https://lvmh.com',
                'employees': 211_000, 'business_summary': 'Luxury goods group.',
            },
            'historical': {'years': ['2023', '2024', '2025'], 'revenue': [86.2e9, 84.7e9, 80.8e9],
                           'gross_profit': [0, 0, 0], 'operating_income': [0, 0, 0],
                           'net_income': [0, 0, 0], 'ebitda': [0, 0, 0],
                           'free_cash_flow': [0, 0, 0]},
            'assumptions': {
                'wacc': 0.0857, 'terminal_growth': 0.025,
                'revenue_growth_rates': five, 'gross_margins': five,
                'ebitda_margins': five, 'operating_margins': five,
            },
            'cost_of_capital': {
                'risk_free_rate': 0.0324, 'equity_risk_premium': 0.055, 'beta': 0.894,
                'cost_of_equity': 0.0816, 'pre_tax_cost_of_debt': 0.0474,
                'tax_rate': 0.3279, 'after_tax_cost_of_debt': 0.0319,
                'equity_weight': 0.8591, 'debt_weight': 0.1409, 'wacc': 0.0746,
            },
            'projections': {
                'revenue': [83.2e9] * 5, 'ebitda': [26.6e9] * 5,
                'fcf': [15.34e9, 14.93e9, 14.35e9, 13.71e9, 12.96e9],
                'ebit': [18.4e9] * 5, 'nopat': [12.4e9] * 5,
            },
            'valuation': {
                'dcf_perpetual': {'pv_fcfs': 83.6e9, 'terminal_value': 176.6e9,
                                  'enterprise_value': 247.3e9, 'equity_value': 232.8e9,
                                  'intrinsic_value_per_share': 467.58},
                'dcf_exit': {'terminal_ev': 200e9, 'enterprise_value': 240e9,
                             'equity_value': 225e9, 'intrinsic_value_per_share': 452.0,
                             'exit_multiple': 12.0},
                'summary': {'dcf_intrinsic': 467.58, 'exit_intrinsic': 452.0,
                            'average_intrinsic': 459.79, 'upside': 0.018,
                            'shares_outstanding': 497_976_118, 'cash': 8.79e9,
                            'debt': 36.73e9, 'net_debt': 27.94e9},
            },
            'news': {
                'summary': {'articles_analyzed': 6, 'overall_sentiment': 'neutral',
                            'confidence_score': 0.6},
                'articles': [], 'catalysts': [], 'risks': [], 'evidence': [],
                'mitigations': [], 'themes': [],
            },
        }

    def _sections(self):
        keys = ('executive_summary', 'company_overview', 'financial_performance',
                'valuation', 'news_analysis', 'investment_thesis', 'recommendation')
        return {k: f"_{k} body_" for k in keys}

    def test_it_assembles_without_raising(self):
        from src.report_agent import integrate_report_sections
        report = integrate_report_sections(self._sections(), self._data())
        assert isinstance(report, str) and len(report) > 500

    def test_the_appendix_quotes_the_rate_the_dcf_used(self):
        """
        It printed assumptions['wacc'] — a different tab's number. On a shipped
        AAPL report that was 8.5% against a model discounting at 11.15%.
        """
        from src.report_agent import integrate_report_sections
        report = integrate_report_sections(self._sections(), self._data())
        assert "| WACC | 7.5%" in report or "| WACC | 7.46%" in report
        assert "8.6%" not in report.split("### C. Key Model Assumptions")[-1]

    def test_it_survives_a_model_with_no_cost_of_capital(self):
        """Older runs and degraded models have no such key."""
        from src.report_agent import integrate_report_sections
        data = self._data()
        data.pop('cost_of_capital')
        report = integrate_report_sections(self._sections(), data)
        assert "| WACC |" in report


class TestTerminalGrowthCannotOutgrowTheCurrency:
    """
    The 2-3% terminal growth band is a dollar band. Once yen and yuan cash
    flows were discounted at their own rates, a 2.5% perpetuity against a 2.3%
    JPY or 1.1% CNY risk-free rate put Toyota at 2.5x its price and Alibaba at
    1.5x: the terminal value, not the business, was doing the valuing. The
    risk-free rate caps the growth rate (Damodaran).
    """

    def _ground(self, currency, country, tg):
        from src.agents.fm.assumption_grounding import ground_assumptions
        return ground_assumptions(
            {"wacc": 0.09, "terminal_growth_rate": tg},
            {"company_data": {"basic_info": {"currency": currency, "country": country},
                              "capital_structure": {"beta": 1.0, "total_debt": 0},
                              "market_data": {"market_cap": 1e12}}},
        )

    def test_yen_growth_is_capped_at_the_yen_risk_free_rate(self):
        a, notes = self._ground("JPY", "Japan", 0.025)
        assert a["terminal_growth_rate"] == pytest.approx(a["capm"]["risk_free_rate"])
        assert a["terminal_growth_rate"] < 0.025
        assert any("capped at the JPY risk-free rate" in n for n in notes)

    def test_yuan_growth_likewise(self):
        a, _ = self._ground("CNY", "China", 0.03)
        assert a["terminal_growth_rate"] == pytest.approx(a["capm"]["risk_free_rate"])
        assert a["terminal_growth_rate"] < 0.02

    def test_a_dollar_perpetuity_is_untouched(self):
        a, notes = self._ground("USD", "United States", 0.025)
        assert a["terminal_growth_rate"] == pytest.approx(0.025)
        assert not any("capped" in n for n in notes)

    def test_a_zero_rate_currency_gets_zero_growth_not_negative(self, monkeypatch):
        monkeypatch.setenv("RISK_FREE_CHF", "0.0")
        a, _ = self._ground("CHF", "Switzerland", 0.02)
        assert a["terminal_growth_rate"] == 0.0

    def test_the_workbook_still_discounts_above_growth(self):
        a, _ = self._ground("JPY", "Japan", 0.025)
        assert a["wacc"] > a["terminal_growth_rate"]


class TestReviewFindings:
    """Each test here is a defect an independent reviewer found in the first cut."""

    def test_an_override_sets_the_risk_free_rate_and_debt_follows_it(self, monkeypatch):
        monkeypatch.setenv("RISK_FREE_INR", "0.05")
        d = risk_free_details("INR")
        assert d["rate"] == pytest.approx(0.05)
        c = capm_components(INDIAN_MIDCAP)
        assert c["pre_tax_cost_of_debt"] == pytest.approx(0.05 + c["domicile_default_spread"] + 0.015)
        assert c["kd_source"].startswith("risk-free 5.00% (RISK_FREE_INR override) + India sovereign spread")

    def test_an_override_with_no_bond_behind_it_is_named_as_such(self, monkeypatch):
        monkeypatch.setenv("RISK_FREE_XXX", "0.05")
        c = capm_components({"basic_info": {"currency": "XXX", "country": "Atlantis"},
                             "capital_structure": {"beta": 1.0, "total_debt": 0},
                             "market_data": {"market_cap": 1e9}, "growth_profitability": {}})
        assert c["kd_source"] == "risk-free 5.00% (RISK_FREE_XXX override) + 1.5% credit spread"
        assert c["pre_tax_cost_of_debt"] == pytest.approx(0.065)

    def test_the_us_proxy_has_the_us_spread_taken_out_like_any_other_us_bond(self):
        d = risk_free_details("XXX")
        us = risk_free_details("USD")
        assert d["proxy"] is True
        assert d["default_spread"] == pytest.approx(us["default_spread"])
        assert d["rate"] == pytest.approx(us["rate"])
        assert "US" in d["default_spread_label"] and "proxy" in d["label"]

    def test_an_aaa_domicile_borrows_over_the_default_free_rate(self):
        c = capm_components(LVMH)                        # no country -> Germany assumed, Aaa
        assert c["kd_source"] == "risk-free 3.36% (default-free) + 1.5% credit spread"
        assert c["pre_tax_cost_of_debt"] == pytest.approx(_SNAPSHOT["EUR"][0] + 0.015)

    def test_an_italian_euro_issuer_borrows_over_the_btp_not_the_bund(self):
        it = capm_components({"basic_info": {"currency": "EUR", "country": "Italy"},
                              "capital_structure": {"beta": 1.0, "total_debt": 1e9},
                              "market_data": {"market_cap": 5e9}, "growth_profitability": {}})
        de = capm_components({"basic_info": {"currency": "EUR", "country": "Germany"},
                              "capital_structure": {"beta": 1.0, "total_debt": 1e9},
                              "market_data": {"market_cap": 5e9}, "growth_profitability": {}})
        assert it["risk_free_rate"] == de["risk_free_rate"]               # same currency, same Rf
        assert it["pre_tax_cost_of_debt"] - de["pre_tax_cost_of_debt"] == pytest.approx(it["domicile_default_spread"])
        assert 0.01 < it["domicile_default_spread"] < 0.03               # Baa2
        assert "Italy sovereign spread" in it["kd_source"]

    def test_a_uk_major_reporting_in_dollars_borrows_over_the_uk_not_the_us(self):
        c = capm_components({"basic_info": {"currency": "USD", "country": "United Kingdom"},
                             "capital_structure": {"beta": 0.7, "total_debt": 8e10},
                             "market_data": {"market_cap": 2e11}, "growth_profitability": {}})
        assert c["domicile"] == "United Kingdom"
        assert c["pre_tax_cost_of_debt"] == pytest.approx(c["risk_free_rate"] + c["domicile_default_spread"] + 0.015)

    def test_no_country_assumes_the_currency_s_sovereign_and_says_so(self):
        eur = capm_components(LVMH)                       # LVMH fixture carries no country
        assert eur["country_risk_premium"] == 0.0
        assert eur["crp_source"].startswith("no country on record — Germany assumed from EUR")
        usd = capm_components(US_MEGACAP)
        assert 0.001 < usd["country_risk_premium"] < 0.005
        assert "United States assumed from USD" in usd["crp_source"]

    def test_a_country_the_table_does_not_carry_also_assumes_the_sovereign(self):
        c = capm_components({"basic_info": {"currency": "GBP", "country": "Atlantis"},
                             "capital_structure": {"beta": 1.0, "total_debt": 0},
                             "market_data": {"market_cap": 1e9}, "growth_profitability": {}})
        assert c["country_risk_premium"] > 0
        assert c["crp_source"].startswith("Atlantis is not in Damodaran's table — United Kingdom assumed from GBP")

    def test_no_country_and_no_known_sovereign_is_one_sentence(self):
        c = capm_components({"basic_info": {"currency": "XXX"},
                             "capital_structure": {"beta": 1.0, "total_debt": 0},
                             "market_data": {"market_cap": 1e9}, "growth_profitability": {}})
        assert c["country_risk_premium"] == 0.0
        assert c["crp_source"] == "no country on record and no sovereign known for XXX — no country premium applied"

    def test_the_wacc_band_follows_the_currency(self):
        yen = capm_components({"basic_info": {"currency": "JPY", "country": "Japan"},
                               "capital_structure": {"beta": 0.7, "total_debt": 3e13},
                               "market_data": {"market_cap": 4e13}, "growth_profitability": {}})
        assert yen["wacc"] < 0.06
        assert yen["wacc_clamped"] is False              # a 5% yen WACC is not a broken input
        assert yen["wacc_band"][0] == pytest.approx(yen["risk_free_rate"] + 0.02)
        _, note = compute_capm_wacc({"basic_info": {"currency": "JPY", "country": "Japan"},
                                     "capital_structure": {"beta": 0.7, "total_debt": 3e13},
                                     "market_data": {"market_cap": 4e13}, "growth_profitability": {}})
        assert "OUTSIDE" not in note

    def test_the_premium_note_matches_its_source_to_the_basis_point(self):
        """0.2% in one half of the cell and 0.23% in the other is a contradiction."""
        import openpyxl
        from src.agents.fm.tabs.tab_assumptions import AssumptionsTabBuilder
        c = capm_components(US_MEGACAP)
        ws = AssumptionsTabBuilder({"capm": c, "terminal_growth_note": "capped at the USD risk-free rate 4.55%"}).create_tab(openpyxl.Workbook())
        note = ws["C24"].value
        assert "country premium 0.23%" in note and "Damodaran's implied base 4.23%" in note and "house assumption" in note
        assert ws["C27"].value == f"[{c['kd_source']}]"
        assert ws["C30"].value == "[capped at the USD risk-free rate 4.55%]"
        assert ws.column_dimensions["C"].width >= 60
        _, chat = compute_capm_wacc(US_MEGACAP)
        assert "CRP 0.23%" in chat

    def test_the_exit_leg_is_capped_at_the_same_growth_as_the_perpetuity(self):
        from src.agents.fm.terminal_value import sustainable_growth_cap, reconcile
        assert sustainable_growth_cap(0.0231) == pytest.approx(0.0231)
        assert sustainable_growth_cap(0.0455) == pytest.approx(0.04)
        assert sustainable_growth_cap(None) == pytest.approx(0.04)
        assert sustainable_growth_cap(-0.01) == pytest.approx(0.04)
        # 6x on 30% conversion at 8.7% implies ~3.5% growth: fine against 4%, not against a 2.31% yen cap.
        loose = reconcile(fcf_terminal=300, ebitda_terminal=1000, wacc=0.0872, terminal_growth=0.0231, exit_multiple=6)
        tight = reconcile(fcf_terminal=300, ebitda_terminal=1000, wacc=0.0872, terminal_growth=0.0231, exit_multiple=6,
                          growth_cap=0.0231)
        assert loose["verdict"] != "growth_not_sustainable"
        assert tight["verdict"] == "growth_not_sustainable" and "2.3%" in tight["note"]

    def test_grounding_defers_the_exit_cap_until_projected_conversion_exists(self):
        from src.agents.fm.assumption_grounding import ground_assumptions
        a, notes = ground_assumptions(
            {"wacc": 0.09, "terminal_growth_rate": 0.025, "exit_multiple": 20.0},
            {"company_data": {"basic_info": {"currency": "JPY", "country": "Japan"},
                              "capital_structure": {"beta": 1.0, "total_debt": 0},
                              "market_data": {"market_cap": 1e12},
                              "valuation_metrics": {"enterprise_to_ebitda": 30.0}},
             "financial_statements": {"cash_flow": {"2025": {"Free Cash Flow": 300.0}},
                                      "income_statement": {"2025": {"EBITDA": 1000.0}}}},
        )
        assert a["sustainable_growth_cap"] == pytest.approx(a["capm"]["risk_free_rate"])
        assert a["exit_multiple"] == pytest.approx(22.0)
        assert any("deferred to projected FY5 cash conversion" in n for n in notes)

    def test_historical_cash_conversion_cannot_rewrite_the_exit_input(self):
        from src.agents.fm.assumption_grounding import ground_assumptions

        base = {
            "company_data": {
                "basic_info": {"currency": "USD", "country": "United States"},
                "capital_structure": {"beta": 1.0, "total_debt": 0},
                "market_data": {"market_cap": 1e12},
                "valuation_metrics": {"enterprise_to_ebitda": 18.0},
            },
            "financial_statements": {
                "income_statement": {"2025": {"EBITDA": 1000.0}},
                "cash_flow": {"2025": {"Free Cash Flow": 100.0}},
            },
        }
        low, _ = ground_assumptions(
            {"wacc": 0.09, "terminal_growth_rate": 0.025}, base)
        base["financial_statements"]["cash_flow"]["2025"]["Free Cash Flow"] = 800.0
        high, _ = ground_assumptions(
            {"wacc": 0.09, "terminal_growth_rate": 0.025}, base)
        assert low["exit_multiple"] == high["exit_multiple"] == pytest.approx(14.4)

    def test_the_report_grid_never_shows_negative_growth(self):
        from src.report_agent import build_sensitivity_grid
        grid = build_sensitivity_grid({}, 0.0031, {"dcf_inputs": {"fcf": [100] * 10, "cash": 0, "debt": 0,
                                                                  "investments": 0, "shares": 10}}, 0.0406)
        header = grid.splitlines()[0]
        assert "-" not in header.replace("| WACC \\ terminal g |", "")
        assert header.count("%") == 5 and "0.3%" in header

    def test_the_workbook_grid_step_shrinks_with_the_base(self):
        from src.agents.fm.tabs.tab_sensitivity import SensitivityTabBuilder
        import inspect
        src = inspect.getsource(SensitivityTabBuilder)
        assert '"=$B$7-IF($B$7<0.01,$B$7/2,0.005)"' in src and '"=$B$7+IF($B$7<0.01,$B$7/2,0.005)"' in src

    def test_a_bank_s_cost_of_equity_is_the_same_build_the_report_prints(self):
        from src.agents.fm.bank_valuation import compute_bank_fair_value
        data = {"valuation_metrics": {"book_value": 100.0}, "growth_profitability": {"return_on_equity": 0.12},
                "market_data": {"current_price": 150.0}, "capital_structure": {"beta": 1.3}}
        old = compute_bank_fair_value(data, 0.025)
        assert old["inputs"]["cost_of_equity"] == pytest.approx(min(0.14, max(0.08, 0.043 + 1.3 * 0.05)))
        capm = {"risk_free_rate": 0.0455, "equity_risk_premium_total": 0.0573, "beta": 1.1}
        new = compute_bank_fair_value(data, 0.025, capm=capm)
        assert new["inputs"]["cost_of_equity"] == pytest.approx(0.0455 + 1.1 * 0.0573)
        assert new["inputs"]["beta"] == pytest.approx(1.1)
        assert "CAPM build" in new["inputs"]["cost_of_equity_source"]
