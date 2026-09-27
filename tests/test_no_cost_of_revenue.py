"""Issuers that never report a cost of revenue must still get a coherent model.

Booking Holdings presents expenses by function: no cost of revenue and no
gross profit in any year. Two defects followed. The Historical tab required a
cost-of-revenue line for a year to count as complete, so no year qualified and
every Booking model failed the historical-period integrity check. And payable
days measured against the absent line were blank, so the forecast carried zero
payables and the first projected year absorbed a $4.35B working-capital outflow
that never happens (payables are 19% of Booking's revenue).
"""

import logging

import openpyxl
import pytest

from src.agents.fm.assumption_grounding import (
    _working_capital_history,
    ground_assumptions,
)
from src.agents.fm.financial_metrics import working_capital_cost_base
from src.agents.fm.formula_evaluator import FormulaEvaluator
from src.agents.fm.tabs.tab_assumptions import (
    AssumptionsTabBuilder,
    source_grounded_assumption_seed,
)
from src.agents.fm.tabs.tab_historical import HistoricalTabBuilder
from src.agents.fm.tabs.tab_projections import ProjectionsTabBuilder
from src.agents.fm.tabs.tab_raw import RawTabBuilder


# Booking Holdings' reported figures, $M (FY2023-FY2025).
REVENUE = {"2025-12-31": 26917.0, "2024-12-31": 23739.0, "2023-12-31": 21365.0}
PAYABLES = {"2025-12-31": 5094.0, "2024-12-31": 3824.0, "2023-12-31": 3374.0}
RECEIVABLES = {"2025-12-31": 3820.0, "2024-12-31": 3199.0, "2023-12-31": 3253.0}


def _booking_statements():
    income, balance, cash = {}, {}, {}
    for period, revenue in REVENUE.items():
        income[period] = {
            "Total Revenue": revenue,
            "Operating Revenue": revenue,
            "Operating Income": revenue * 0.34,
            "Net Income": revenue * 0.22,
            "Operating Expense": revenue * 0.66,
        }
        balance[period] = {
            "Accounts Receivable": RECEIVABLES[period],
            "Accounts Payable": PAYABLES[period],
            "Cash And Cash Equivalents": 16000.0,
            "Total Debt": 17000.0,
        }
        cash[period] = {
            "Operating Cash Flow": revenue * 0.33,
            "Capital Expenditure": -revenue * 0.03,
            "Free Cash Flow": revenue * 0.30,
        }
    return {"income_statement": income, "balance_sheet": balance, "cash_flow": cash}


def _booking_grounding_data():
    return {
        "company_data": {
            "basic_info": {"currency": "USD", "country": "United States"},
            "market_data": {"market_cap": 123_000.0, "shares_outstanding": 751.0},
            "capital_structure": {"beta": 1.1, "total_debt": 17000.0},
            # Yahoo's info-level gross margin exists even though no statement
            # reports a cost of revenue; it must not become the payables base.
            "growth_profitability": {
                "operating_margins": 0.34,
                "ebitda_margins": 0.37,
                "gross_margins": 0.87,
            },
            "valuation_metrics": {},
        },
        "financial_statements": _booking_statements(),
    }


def _days(amount, base):
    return amount / base * 365.0


# --- the shared rule ---------------------------------------------------------

def test_issuer_without_any_cost_structure_measures_days_against_revenue():
    assert working_capital_cost_base(_booking_statements()) == "revenue"


@pytest.mark.parametrize("fields", [
    {"Cost Of Revenue": 60.0},
    {"Reconciled Cost Of Revenue": 60.0},
    # A gross profit below revenue implies a cost of revenue.
    {"Gross Profit": 40.0},
])
def test_any_reported_cost_structure_keeps_the_cost_of_revenue_base(fields):
    statements = _booking_statements()
    statements["income_statement"]["2023-12-31"].update(fields)

    assert working_capital_cost_base(statements) == "cost_of_revenue"


def test_empty_statements_keep_the_default_base():
    assert working_capital_cost_base({}) == "cost_of_revenue"
    assert working_capital_cost_base({"income_statement": {}}) == "cost_of_revenue"


# --- grounding ---------------------------------------------------------------

def test_payable_days_reproduce_the_reported_payables_to_revenue_ratio():
    history = _working_capital_history(
        {"financial_statements": _booking_statements()}, "dpo_days"
    )

    assert history == pytest.approx([
        _days(PAYABLES[p], REVENUE[p])
        for p in ("2025-12-31", "2024-12-31", "2023-12-31")
    ])


def test_grounding_carries_the_revenue_base_to_the_workbook():
    grounded, notes = ground_assumptions({
        "wacc": 0.09, "terminal_growth_rate": 0.025,
        "dso_days": [None] * 5, "dio_days": [None] * 5, "dpo_days": [None] * 5,
    }, _booking_grounding_data())

    assert grounded["working_capital_cost_base"] == "revenue"
    latest = _days(PAYABLES["2025-12-31"], REVENUE["2025-12-31"])
    median = _days(PAYABLES["2024-12-31"], REVENUE["2024-12-31"])
    assert grounded["dpo_days"][0] == pytest.approx(latest)      # 69.1 days
    assert grounded["dpo_days"][-1] == pytest.approx(median)     # 58.8 days
    assert any("measured against revenue" in note for note in notes)


def test_grounding_leaves_a_cost_of_revenue_issuer_unchanged():
    data = _booking_grounding_data()
    for period, revenue in REVENUE.items():
        data["financial_statements"]["income_statement"][period][
            "Cost Of Revenue"] = revenue * 0.5

    grounded, notes = ground_assumptions({
        "wacc": 0.09, "terminal_growth_rate": 0.025,
        "dso_days": [None] * 5, "dio_days": [None] * 5, "dpo_days": [None] * 5,
    }, data)

    assert grounded["working_capital_cost_base"] == "cost_of_revenue"
    assert grounded["dpo_days"][0] == pytest.approx(
        _days(PAYABLES["2025-12-31"], REVENUE["2025-12-31"] * 0.5)
    )
    assert not any("measured against revenue" in note for note in notes)


def test_fallback_seed_uses_the_same_base():
    seed = source_grounded_assumption_seed({
        "financial_statements": _booking_statements(),
        "modeling_metrics": {
            "historical_growth_rates": {"revenue_growth": {"cagr_3y": 0.12}}
        },
    })

    assert seed["working_capital_cost_base"] == "revenue"
    assert seed["dpo_days"][0] == pytest.approx(
        _days(PAYABLES["2025-12-31"], REVENUE["2025-12-31"])
    )


# --- workbook ----------------------------------------------------------------

def _historical_cells(statements, cost_base=None):
    workbook = openpyxl.Workbook()
    raw = RawTabBuilder()
    raw.add_data_from_json({"financial_statements": statements})
    raw.create_tab(workbook)
    HistoricalTabBuilder(working_capital_cost_base=cost_base).create_tab(workbook)
    evaluator = FormulaEvaluator(workbook)
    evaluator.set_logger(logging.getLogger("test_no_cost_of_revenue"))
    return evaluator.evaluate_all_tabs()["Historical"]["cells"]


def test_every_booking_year_is_a_complete_historical_period():
    cells = _historical_cells(_booking_statements(), "revenue")

    assert [cells.get(f"(1, {col})") for col in (4, 5, 6)] == [2023, 2024, 2025]
    assert cells["(41, 6)"] == pytest.approx(
        round(_days(PAYABLES["2025-12-31"], REVENUE["2025-12-31"]), 2)
    )
    # The accounting identities the integrity check reads still hold.
    assert cells["(5, 6)"] == pytest.approx(REVENUE["2025-12-31"])
    assert cells["(45, 6)"] == pytest.approx(0.0, abs=1e-6)


def test_a_sparse_stub_year_still_needs_cost_of_revenue():
    statements = _booking_statements()
    for period, revenue in REVENUE.items():
        statements["income_statement"][period]["Cost Of Revenue"] = revenue * 0.5
    # Yahoo's oldest column is usually a stub without cost of revenue.
    statements["income_statement"]["2022-12-31"] = {
        "Total Revenue": 17090.0, "Operating Income": 5102.0, "Net Income": 3058.0,
    }
    statements["balance_sheet"]["2022-12-31"] = {"Cash And Cash Equivalents": 12000.0}
    statements["cash_flow"]["2022-12-31"] = {"Operating Cash Flow": 6000.0}

    cells = _historical_cells(statements)

    assert [cells.get(f"(1, {col})") for col in (4, 5, 6)] == [2023, 2024, 2025]
    assert cells.get("(1, 3)") in (None, "")


def test_assumptions_tab_divides_by_the_grounded_base():
    workbook = openpyxl.Workbook()
    AssumptionsTabBuilder({"working_capital_cost_base": "revenue"}).create_tab(workbook)
    sheet = workbook["Assumptions"]
    assert '"Total Revenue"' in sheet["B15"].value
    assert '"Cost Of Revenue"' not in sheet["B15"].value
    assert '"Total Revenue"' in sheet["B14"].value
    assert "no cost of revenue" in sheet["H15"].value
    # Readers find these rows by label; the label must not change.
    assert sheet["A15"].value == "DPO (Days)"
    # Not reported is not zero: the latest gross margin stays blank.
    assert sheet["B9"].value is None
    assert "Not reported" in sheet["H9"].value
    assert sheet["C9"].value.startswith("=IF(IFERROR(Model_Inputs!B5")

    workbook = openpyxl.Workbook()
    AssumptionsTabBuilder({}).create_tab(workbook)
    sheet = workbook["Assumptions"]
    assert '"Cost Of Revenue"' in sheet["B15"].value
    assert sheet["H15"].value is None
    assert '"Gross Profit"' in sheet["B9"].value
    assert sheet["H9"].value is None


def test_projected_payables_and_inventory_use_the_grounded_base():
    workbook = openpyxl.Workbook()
    sheet = ProjectionsTabBuilder(working_capital_cost_base="revenue").create_tab(workbook)
    assert sheet["B16"].value == "=B3/365*Assumptions!C15"
    assert sheet["F15"].value == "=F3/365*Assumptions!G14"

    workbook = openpyxl.Workbook()
    sheet = ProjectionsTabBuilder().create_tab(workbook)
    assert sheet["B16"].value == "=B4/365*Assumptions!C15"
    assert sheet["F15"].value == "=F4/365*Assumptions!G14"
