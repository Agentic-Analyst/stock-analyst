import pytest
import openpyxl
from datetime import datetime, timedelta, timezone

from src.agents.fm.bank_valuation import (
    assess_bank_publication,
    build_bank_valuation_override,
    compute_bank_fair_value,
    is_balance_sheet_financial,
)
from src.report_agent import (
    enforce_valuation_publication_boundary,
    generate_section_valuation,
)
from src.agents.fm.formula_evaluator import FormulaEvaluator, formula_integrity
from src.agents.fm.financial_model_builder import FinancialModelBuilder
from src.agents.fm.tabs.tab_bank_valuation import (
    BankValuationTabBuilder,
    apply_bank_summary,
)
from src.agents.fm.tabs.tab_summary import SummaryTabBuilder


def _bank_financials():
    balance = {
        "2026-06-30": {
            "Common Stock Equity": 1_100.0,
            "Ordinary Shares Number": 10.0,
            "Total Assets": 10_000.0,
        },
        "2025-06-30": {"Common Stock Equity": 1_000.0, "Total Assets": 9_000.0},
        "2024-06-30": {"Common Stock Equity": 900.0, "Total Assets": 8_000.0},
        "2023-06-30": {"Common Stock Equity": 800.0, "Total Assets": 7_000.0},
    }
    income = {
        "2026-06-30": {"Net Income Common Stockholders": 210.0, "Total Revenue": 500.0},
        "2025-06-30": {"Net Income Common Stockholders": 190.0, "Total Revenue": 450.0},
        "2024-06-30": {"Net Income Common Stockholders": 170.0, "Total Revenue": 400.0},
    }
    peer_rows = [
        {"symbol": "A", "price_to_book": 1.2, "return_on_equity_ttm_pct": 10.0},
        {"symbol": "B", "price_to_book": 1.6, "return_on_equity_ttm_pct": 12.0},
        {"symbol": "C", "price_to_book": 2.0, "return_on_equity_ttm_pct": 14.0},
    ]
    return {
        "company_data": {
            "basic_info": {
                "quote_type": "EQUITY",
                "currency": "USD",
                "sector": "Financial Services",
                "industry": "Banks - Diversified",
            },
            "valuation_metrics": {"book_value": 90.0},
            "growth_profitability": {"return_on_equity": .30},
            "market_data": {"current_price": 250.0, "market_cap": 1e11},
            "capital_structure": {"beta": 1.0},
            "forward_guidance": {
                "target_mean_price": 262.5,
                "number_of_analyst_opinions": 20,
            },
            "analyst_consensus": {
                "recommendation": {"label": "buy", "total": 20},
            },
        },
        "financial_statements": {
            "income_statement": income,
            "balance_sheet": balance,
            "cash_flow": {
                "2026-06-30": {"Operating Cash Flow": 1.0},
                "2025-06-30": {"Operating Cash Flow": 1.0},
                "2024-06-30": {"Operating Cash Flow": 1.0},
            },
        },
        "quarterly_financial_statements": {
            "balance_sheet": {
                "2026-06-30": {
                    "Common Stock Equity": 1_000.0,
                    "Ordinary Shares Number": 10.0,
                },
            },
        },
        "modeling_metrics": {"financial_ratios": {"financial_profile": {
            "interest_income_to_revenue": 1.2,
        }}},
        "industry_data": {"peer_comps": {
            "status": "ready",
            "selected_method": "price_to_book",
            "included_in_blended_value": True,
            "size_screen_applied": True,
            "fundamental_screen_applied": True,
            "observations": peer_rows,
        }},
    }


def test_bank_model_uses_reported_bvps_normalized_roe_and_adjusted_peers():
    result = build_bank_valuation_override(
        _bank_financials(),
        terminal_growth=.02,
        capm={"risk_free_rate": .04, "equity_risk_premium_total": .06, "beta": 1.0},
    )

    inputs = result["bank_inputs"]
    assert inputs["bvps"] == 100.0
    assert inputs["roe"] == pytest.approx(.20)
    assert inputs["return_on_equity_source"] == "median historical common ROE (3 periods)"
    assert inputs["intrinsic_fair_value"] == 225.0
    # The peer set is trailing, so the subject's provider trailing ROE (30%)
    # is used instead of its three-year normalized ROE (20%).
    assert inputs["peer_subject_return_on_equity"] == pytest.approx(.30)
    assert inputs["peer_implied_price_to_book"] == pytest.approx(4.48)
    assert inputs["peer_fair_value"] == 448.0
    assert result["fair_value"] == 336.5


def test_bank_forward_eps_is_a_fresh_well_covered_first_class_scenario():
    financials = _bank_financials()
    financials["company_data"]["growth_profitability"]["return_on_equity"] = .20
    financials["external_expectations"] = {
        "captured_at": datetime.now(timezone.utc).isoformat(),
        "currency": "USD",
        "forward_estimates": [
            {"period": "0y", "eps": 22.0, "eps_analyst_count": 12},
            {"period": "+1y", "eps": 24.0, "eps_analyst_count": 14},
        ],
        "price_target": {},
        "recommendations": {},
    }

    result = build_bank_valuation_override(
        financials,
        terminal_growth=.02,
        capm={"risk_free_rate": .04, "equity_risk_premium_total": .06, "beta": 1.0},
    )

    inputs = result["bank_inputs"]
    assert inputs["forward_consensus_roe_status"] == "ready"
    assert inputs["forward_consensus_roe"] == pytest.approx(.23)
    assert inputs["forward_consensus_justified_pb"] == pytest.approx(2.625)
    assert result["forward_consensus_fair_value"] == 262.5
    assert result["reliability_legs"]["forward_consensus_roe_scenario"] == 262.5
    assert result["fair_value"] == 258.5


def test_stale_forward_eps_cannot_enter_bank_value():
    financials = _bank_financials()
    financials["external_expectations"] = {
        "captured_at": (datetime.now(timezone.utc) - timedelta(days=90)).isoformat(),
        "currency": "USD",
        "forward_estimates": [
            {"period": "0y", "eps": 22.0, "eps_analyst_count": 12},
            {"period": "+1y", "eps": 24.0, "eps_analyst_count": 14},
        ],
        "price_target": {},
        "recommendations": {},
    }

    result = build_bank_valuation_override(
        financials,
        terminal_growth=.02,
        capm={"risk_free_rate": .04, "equity_risk_premium_total": .06, "beta": 1.0},
    )

    assert result["forward_consensus_fair_value"] is None
    assert result["bank_inputs"]["forward_consensus_roe_status"] == "unavailable"
    assert "older than 30 days" in result["bank_inputs"][
        "forward_consensus_roe_reason"
    ]


def test_bank_input_clamps_are_visible_publication_blockers():
    financials = _bank_financials()
    # Remove statement-derived ROE so the intentionally extreme provider ROE
    # reaches the formula and trips the P/B safety boundary.
    financials["financial_statements"]["income_statement"] = {
        "2025-12-31": {"Total Revenue": 500.0},
        "2024-12-31": {"Total Revenue": 450.0},
    }
    financials["industry_data"] = {}
    override = build_bank_valuation_override(
        financials,
        terminal_growth=.02,
        capm={"risk_free_rate": .04, "equity_risk_premium_total": .05, "beta": 1.0},
    )
    assert override["bank_inputs"]["price_to_book_clamped"] is True
    assert override["bank_inputs"]["input_boundary_triggered"] is True


def test_bank_analyst_conflict_withholds_directional_call():
    financials = _bank_financials()
    data = {
        "company_overview": {
            "current_price": 250.0,
            "market_cap": 1e11,
            "currency": "USD",
            "listing_currency": "USD",
            "target_mean_price": 262.5,
            "num_analysts": 20,
            "analyst_consensus": {
                "recommendation": {"label": "buy", "total": 20},
            },
        },
        "valuation": {
            "bank": {
                "fair_value": 200.0,
                "intrinsic_fair_value": 200.0,
                "peer_fair_value": None,
                "inputs": {},
            },
            "summary": {"average_intrinsic": 200.0, "upside": -.20},
        },
    }

    result = enforce_valuation_publication_boundary(data, financials)

    reliability = result["valuation"]["reliability"]
    assert reliability["point_estimate_withheld"] is True
    assert "20-analyst target benchmark" in reliability["withheld_reason"]
    assert "does not corroborate" in reliability["withheld_reason"]


def test_bank_rating_label_cannot_validate_large_single_method_target():
    result = assess_bank_publication(
        fair_value=150.0,
        intrinsic_fair_value=150.0,
        peer_fair_value=None,
        current_price=100.0,
        analyst_target=None,
        analyst_count=0,
        analyst_rating="strong_buy",
        analyst_rating_count=30,
    )

    assert result["point_estimate_withheld"] is True
    assert "does not corroborate" in result["publication_withheld_reason"]


def test_undated_bank_target_cannot_corroborate_large_single_method_target():
    result = assess_bank_publication(
        fair_value=150.0,
        intrinsic_fair_value=150.0,
        peer_fair_value=None,
        current_price=100.0,
        analyst_target=140.0,
        analyst_count=30,
        analyst_target_corroboration_qualified=False,
    )

    assert result["point_estimate_withheld"] is True


def test_directional_single_method_bank_call_needs_numeric_corroboration():
    result = assess_bank_publication(
        fair_value=120.0,
        intrinsic_fair_value=120.0,
        peer_fair_value=None,
        current_price=100.0,
    )

    assert result["point_estimate_withheld"] is True


def test_stale_bank_target_cannot_veto_a_two_method_result():
    result = assess_bank_publication(
        fair_value=120.0,
        intrinsic_fair_value=118.0,
        peer_fair_value=122.0,
        current_price=100.0,
        analyst_target_evidence={
            "stale_provider": {
                "mean": 80.0,
                "analyst_count": 20,
                "qualified_for_contradiction": False,
                "qualified_for_corroboration": False,
            },
        },
    )

    assert result["point_estimate_withheld"] is False


def test_secondary_bank_analyst_source_can_block_a_conflicting_call():
    result = assess_bank_publication(
        fair_value=120.0,
        intrinsic_fair_value=118.0,
        peer_fair_value=122.0,
        current_price=100.0,
        analyst_target_evidence={
            "primary": {
                "mean": 125.0,
                "analyst_count": 20,
                "qualified_for_contradiction": True,
                "qualified_for_corroboration": True,
            },
            "secondary": {
                "mean": 96.0,
                "analyst_count": 18,
                "qualified_for_contradiction": True,
                "qualified_for_corroboration": False,
            },
        },
    )

    assert result["point_estimate_withheld"] is True
    assert "secondary 18-analyst target benchmark" in result[
        "publication_withheld_reason"
    ]


def test_bank_workbook_exposes_the_selected_method_not_industrial_dcf():
    financials = _bank_financials()
    override = build_bank_valuation_override(
        financials,
        terminal_growth=.02,
        capm={"risk_free_rate": .04, "equity_risk_premium_total": .06, "beta": 1.0},
    )
    workbook = openpyxl.Workbook()
    workbook.active.title = "Summary"
    SummaryTabBuilder(currency="USD").create_tab(workbook)
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])
    BankValuationTabBuilder(override, financials).create_tab(workbook)
    apply_bank_summary(workbook)
    evaluator = FormulaEvaluator(workbook)
    evaluator.TAB_ORDER = ["Bank Valuation", "Summary"]
    result = evaluator.evaluate_all_tabs()

    summary = result["Summary"]["cells"]
    bank = result["Bank Valuation"]["cells"]
    assert summary["(1, 1)"] == "SUMMARY — BALANCE-SHEET FINANCIAL VALUATION"
    assert summary["(18, 2)"] == 225.0
    assert summary["(22, 2)"] == 448.0
    assert summary["(26, 2)"] == 336.5
    assert "Industrial free-cash-flow DCF is not applicable" in summary["(2, 1)"]
    # The workbook exposes the normalization that turns each raw peer P/B into
    # a subject-ROE-implied P/B; the peer result is no longer an opaque number.
    assert bank["(32, 4)"] == pytest.approx(15.0)
    assert bank["(32, 5)"] == pytest.approx(4.2)
    assert workbook["Summary"].max_row == 41
    assert formula_integrity(result)["status"] == "ready"


def test_summary_uses_listing_financial_currency_not_hardcoded_dollars():
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    summary = SummaryTabBuilder(currency="EUR").create_tab(workbook)

    assert "€" in summary["B9"].number_format
    assert "$" not in summary["B9"].number_format
    assert "€" in summary["B13"].number_format


def test_bank_classification_does_not_depend_on_complete_valuation_inputs():
    financials = _bank_financials()
    financials["company_data"]["valuation_metrics"] = {}
    financials["company_data"]["growth_profitability"] = {}
    financials["financial_statements"]["balance_sheet"] = {}
    financials["financial_statements"]["income_statement"] = {
        "2026-06-30": {
            "Total Revenue": 500.0,
            "Interest Income": 600.0,
        }
    }

    assert is_balance_sheet_financial(financials) is True
    assert build_bank_valuation_override(financials) is None


def test_malformed_optional_bank_inputs_cannot_crash_or_distort_the_method():
    company = _bank_financials()["company_data"]
    company["capital_structure"]["beta"] = "not-a-number"
    result = compute_bank_fair_value(
        company,
        capm={
            "risk_free_rate": "also-bad",
            "equity_risk_premium_total": float("nan"),
            "beta": True,
        },
    )

    assert result is not None
    assert result["inputs"]["beta"] == 1.0
    assert "no CAPM build available" in result["inputs"]["cost_of_equity_source"]


def test_missing_bank_inputs_suppress_legacy_industrial_dcf_everywhere():
    financials = _bank_financials()
    financials["company_data"]["valuation_metrics"] = {}
    financials["company_data"]["growth_profitability"] = {}
    financials["financial_statements"]["balance_sheet"] = {}
    financials["financial_statements"]["income_statement"] = {
        "2026-06-30": {
            "Total Revenue": 500.0,
            "Interest Income": 600.0,
        }
    }
    data = {
        "company_overview": {"current_price": 250.0, "currency": "USD"},
        "valuation": {
            "dcf_perpetual": {
                "intrinsic_value_per_share": 99.0,
                "pv_fcfs": 1.0, "terminal_value": 2.0,
                "enterprise_value": 3.0, "equity_value": 4.0,
            },
            "dcf_exit": {
                "intrinsic_value_per_share": 101.0,
                "exit_multiple": 10.0, "terminal_ev": 2.0,
                "enterprise_value": 3.0, "equity_value": 4.0,
            },
            "summary": {
                "dcf_intrinsic": 99.0,
                "exit_intrinsic": 101.0,
                "comps_intrinsic": None,
                "average_intrinsic": 100.0,
                "upside": -.60,
            },
        },
        "assumptions": {
            "terminal_growth": .02,
            "revenue_growth_rates": [.01] * 5,
            "ebitda_margins": [.01] * 5,
            "wacc": .10,
        },
        "projections": {
            "revenue": [1.0] * 5,
            "ebitda": [1.0] * 5,
            "fcf": [1.0] * 5,
        },
        "peer_comps": {},
        "external_expectations": {},
    }

    result = enforce_valuation_publication_boundary(data, financials)

    reliability = result["valuation"]["reliability"]
    assert reliability["point_estimate_withheld"] is True
    assert reliability["method_suitability"]["primary_method"] == "justified_pb_roe"
    assert result["valuation"]["summary"]["average_intrinsic"] is None
    assert result["valuation"]["dcf_perpetual"]["intrinsic_value_per_share"] is None
    assert result["valuation"]["inapplicable_industrial_dcf"][
        "workbook_headline_value"
    ] == 100.0

    section, _ = generate_section_valuation(result, llm=None)
    assert "| Bank valuation inputs | Unavailable or insufficient |" in section
    assert "Industrial FCF DCF | Not applicable and suppressed" in section
    assert "DCF Valuation — Perpetual Growth Method" not in section


def test_workbook_builder_refuses_dcf_when_selected_bank_method_disappears(
    monkeypatch,
):
    class Logger:
        def info(self, _message):
            pass

    monkeypatch.setattr(
        "src.valuation_methodology.assess_valuation_methodology",
        lambda _data: {
            "primary_method": "justified_pb_roe",
            "publication_allowed": True,
        },
    )
    monkeypatch.setattr(
        "src.agents.fm.financial_model_builder.source_grounded_assumption_seed",
        lambda _data: {"terminal_growth_rate": .02, "capm": {},
                       "revenue_growth_rates": []},
    )
    monkeypatch.setattr(
        "src.agents.fm.assumption_grounding.ground_assumptions",
        lambda assumptions, _data: (assumptions, []),
    )
    monkeypatch.setattr(
        "src.agents.fm.bank_valuation.build_bank_valuation_override",
        lambda *_args, **_kwargs: None,
    )
    builder = FinancialModelBuilder("JPM", Logger())
    builder.json_data = {}

    with pytest.raises(ValueError, match="refusing to build an industrial"):
        builder.build_model()
