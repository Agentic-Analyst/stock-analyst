"""Every surface states the same confidence alert, and none drops it.

Meta, 2026-09-27: a sound DCF said SELL at about half the market price while
35 analysts' mean target sat above the price. The engine used to withhold the
answer ("NOT RATED"). It now publishes the answer and says, wherever the answer
appears, how far it stands from the Street: the report's first page, its
valuation and recommendation sections, the workbook's status cell, the chat
answer (templated or free prose, English or Chinese), the tool results the
prose model reads, and the findings card.
"""
import os
import sys
import types
from types import SimpleNamespace

import pytest

_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "src"))
sys.path.insert(0, os.path.dirname(__file__))

from src.agents.supervisor import supervisor_agent  # noqa: E402
from src.confidence_alert import (  # noqa: E402
    alert_body, alert_sentence, alert_sentence_zh, build_street_alert,
)
from src.recommendation_engine import RecommendationEngineV3  # noqa: E402
from agents.findings import extract_findings  # noqa: E402
from test_guarded_answer_keeps_analysis import _Provider, _analysis, _chat_agent  # noqa: E402
from test_report_recommendation import FIXED_NUMBERS, LLM_RESPONSE  # noqa: E402
from test_run_audit_fixes import _valuation_data  # noqa: E402

ALERT = build_street_alert(
    -0.49,
    targets=[{"gap": 0.15, "count": 35, "source": "yahoo_finance"}],
    detail=(
        "The DCF-only estimate is -49% from the market for a mega-cap, but no "
        "independent market-comps valuation qualified and the available "
        "well-covered analyst-target evidence does not corroborate both the "
        "direction and material magnitude of that gap. External benchmarks: "
        "yahoo_finance 35-analyst target benchmark is +15%."
    ),
)
SENTENCE = (
    "Low confidence: VYNN's fair value is 49% below the market price, while the "
    "mean target of 35 analysts is 15% above it. That is a large gap between "
    "VYNN and the Street, so treat this as VYNN's own view and weigh both."
)
BODY = SENTENCE[len("Low confidence: "):]


def test_the_fixture_is_the_sentence_under_test():
    assert alert_sentence(ALERT) == SENTENCE and alert_body(ALERT) == BODY


# --------------------------------------------------------------------------
# The report
# --------------------------------------------------------------------------

def _flagged_report_data():
    from src.report_agent import apply_valuation_override
    data = apply_valuation_override(_valuation_data(comps=None), {
        "valuation_method": "dcf",
        "perpetual_price": 98.34,
        "exit_multiple_price": 89.59,
        "comps_price": None,
        "dispersion_band": "single-method",
        "dispersion_ratio": 1.10,
        "confidence_alert": ALERT,
    })
    data["news"] = {
        "summary": {"overall_sentiment": "neutral"},
        "catalysts": [], "risks": [], "freshness": {"status": "unavailable"},
    }
    return data


def _must_not_call(*_args, **_kwargs):
    raise AssertionError("a flagged valuation is described in code, not by a prose model")


def test_the_override_keeps_the_alert_only_on_a_published_valuation():
    from src.report_agent import apply_valuation_override
    reliability = _flagged_report_data()["valuation"]["reliability"]
    assert reliability["point_estimate_withheld"] is False
    assert reliability["confidence_alert"] == ALERT

    blocked = apply_valuation_override(_valuation_data(comps=None), {
        "valuation_method": "dcf", "perpetual_price": 98.34, "exit_multiple_price": 89.59,
        "dispersion_band": "single-method", "point_estimate_withheld": True,
        "publication_withheld_reason": "Latest annual financial period is 270 days old.",
        "confidence_alert": ALERT,
    })["valuation"]["reliability"]
    assert blocked["point_estimate_withheld"] is True
    assert blocked["confidence_alert"] is None


def test_the_valuation_section_shows_the_value_and_its_confidence_together():
    from src.report_agent import generate_section_valuation
    text, cost = generate_section_valuation(_flagged_report_data(), _must_not_call)
    assert cost == 0.0
    # The fair value is published ...
    assert "| **DCF Fair Value** |" in text
    assert "Single Fair Value" not in text and "Range only" not in text
    # ... and the same table says how much weight it carries.
    assert "| **Confidence** | **Low confidence: far from analyst consensus** |" in text
    assert f"**Confidence alert.** {BODY}" in text
    # The alert follows the value in the commentary; it does not replace it.
    assert text.index("The publishable model value is") < text.index("**Confidence alert.**")


def test_a_corroborated_valuation_has_no_confidence_row():
    from src.report_agent import apply_valuation_override, generate_section_valuation
    data = apply_valuation_override(_valuation_data(comps=None), {
        "valuation_method": "dcf", "perpetual_price": 98.34, "exit_multiple_price": 89.59,
        "comps_price": None, "dispersion_band": "single-method", "dispersion_ratio": 1.10,
    })
    text, _ = generate_section_valuation(data, _must_not_call)
    assert "| **Confidence** |" not in text and "Confidence alert" not in text


def test_the_status_block_gives_the_alert_with_the_engine_s_reasoning():
    from src.report_agent import valuation_publication_status
    status = valuation_publication_status(_flagged_report_data())
    lines = status.splitlines()
    assert lines[0] == "### Confidence Alert"
    assert lines[2] == f"**Low confidence.** {BODY}"
    assert ALERT["detail"] in status
    assert "are not blended into them" in status
    # A flagged run is published: nothing here says range-only.
    assert "Range only" not in status and "Supported Valuation Range" not in status


def test_a_corroborated_valuation_has_no_status_block():
    from src.report_agent import apply_valuation_override, valuation_publication_status
    data = apply_valuation_override(_valuation_data(comps=None), {
        "valuation_method": "dcf", "perpetual_price": 98.34, "exit_multiple_price": 89.59,
        "dispersion_band": "single-method",
    })
    assert valuation_publication_status(data) == ""


def test_the_thesis_states_the_alert_after_the_published_value():
    from src.report_agent import generate_section_investment_thesis
    text, cost = generate_section_investment_thesis(_flagged_report_data(), _must_not_call)
    assert cost == 0.0
    assert "The model publishes an intrinsic value of" in text
    assert text.index("The model publishes an intrinsic value of") < text.index(SENTENCE)
    assert "Evidence Required Before a Directional Call" not in text


# --------------------------------------------------------------------------
# The recommendation section
# --------------------------------------------------------------------------

@pytest.fixture
def engine():
    eng = RecommendationEngineV3(sector="default")
    eng._ccy = "$"
    return eng


def _flagged_numbers():
    return dict(FIXED_NUMBERS, rating_available=True, rating_confidence="low",
                confidence_alert=ALERT, confidence_alert_text=SENTENCE)


def _header(text):
    return text.split("\n### Investment Thesis")[0].split("\n### Source-Grounded")[0]


def test_the_formatted_recommendation_prints_the_alert_under_the_rating(engine):
    out = engine._format_final_output(LLM_RESPONSE, _flagged_numbers(), {})
    lines = out.splitlines()
    assert lines[0] == "### Investment Rating: SELL"
    assert lines[1] == "**Rating Confidence**: Low"
    assert lines[2] == f"**Confidence Alert**: {BODY}"
    # The rating and target still stand, exactly as the calculator computed them.
    assert "**12-Month Price Target**: $402.32" in out


def test_both_fallbacks_print_the_alert_too(engine):
    minimal = engine._minimal_recommendation(_flagged_numbers())
    assert minimal.splitlines()[:3] == [
        "### Investment Rating: SELL",
        "**Rating Confidence**: Low",
        f"**Confidence Alert**: {BODY}",
    ]
    safe = engine._evidence_safe_recommendation(_flagged_numbers(), {"evidence": []}, {})
    assert safe.splitlines()[:3] == minimal.splitlines()[:3]


def test_a_corroborated_rating_prints_no_alert(engine):
    for out in (
        engine._format_final_output(LLM_RESPONSE, FIXED_NUMBERS, {}),
        engine._minimal_recommendation(FIXED_NUMBERS),
    ):
        assert "Confidence Alert" not in out


def _unrated_numbers(low=None, high=None):
    numbers = dict(_flagged_numbers(), rating="NOT RATED", rating_available=False,
                   rating_withheld_reason="The valuation methods span more than 1.8x.")
    numbers["inputs"] = dict(FIXED_NUMBERS["inputs"], valuation_reliability={
        "band": "wide", "range_low": low, "range_high": high, "point_estimate_withheld": True,
    })
    return numbers


def test_an_alert_is_never_printed_without_a_rating(engine):
    unrated = _unrated_numbers(150.0, 290.0)
    for out in (
        engine._format_final_output(LLM_RESPONSE, unrated, {}),
        engine._minimal_recommendation(unrated),
        engine._evidence_safe_recommendation(unrated, {"evidence": []}, {}),
    ):
        assert out.splitlines()[0] == "### Investment View: Range Only"
        assert "Confidence Alert" not in out and "NOT RATED" not in out


@pytest.mark.parametrize("low,high,heading,status,value_line,statement", [
    (150.0, 290.0, "Range Only", "Range only", "**Supported Valuation Range**: $150.00 – $290.00",
     "VYNN shows a range here rather than a single fair value, rating or price target."),
    (264.73, 264.73, "Scenario Estimate Only", "Scenario estimate only",
     "**Supported Scenario Estimate**: $264.73",
     "VYNN shows one scenario estimate here, not a fair value, rating or price target."),
    (None, None, "No Fair Value", "None", None,
     "VYNN states no fair value, rating or price target here."),
])
def test_the_words_follow_what_the_methods_support(engine, low, high, heading, status, value_line, statement):
    # One surviving method is not a range, and nothing is not a range either.
    unrated = _unrated_numbers(low, high)
    for out in (
        engine._format_final_output(LLM_RESPONSE, unrated, {}),
        engine._minimal_recommendation(unrated),
        engine._evidence_safe_recommendation(unrated, {"evidence": []}, {}),
    ):
        assert out.splitlines()[0] == f"### Investment View: {heading}"
        assert f"**Single Fair Value**: {status}" in out
        assert f"**{statement}**" in out
        if value_line:
            assert value_line in out
        else:
            assert "Supported Valuation Range" not in out and "Supported Scenario Estimate" not in out
        if heading != "Range Only":
            assert "a range" not in out.split("### Source-Grounded")[0].split("> **")[0].lower()


def test_the_explainer_is_told_to_state_the_disagreement_and_not_soften_it(engine):
    prompt = engine._build_explainer_prompt(
        _flagged_numbers(), {"evidence": []}, {}, {}, citations_enabled=False)
    assert "PUBLISHED WITH A CONFIDENCE ALERT" in prompt
    assert f"Alert: {BODY}" in prompt
    assert "Never describe the rating as supported by, consistent with, or confirmed by" in prompt
    plain = engine._build_explainer_prompt(
        dict(FIXED_NUMBERS, rating_available=True), {"evidence": []}, {}, {},
        citations_enabled=False)
    assert "CONFIDENCE ALERT" not in plain


# --------------------------------------------------------------------------
# The workbook
# --------------------------------------------------------------------------

def _builder():
    import openpyxl

    from src.agents.fm.financial_model_builder import FinancialModelBuilder
    from src.agents.fm.tabs.tab_sensitivity import SensitivityTabBuilder
    from src.agents.fm.tabs.tab_summary import SummaryTabBuilder

    class Logger:
        def info(self, _message):
            pass

    builder = FinancialModelBuilder("META", Logger())
    builder.workbook = openpyxl.Workbook()
    builder.workbook.remove(builder.workbook.active)
    builder.summary_builder = SummaryTabBuilder(currency="USD")
    builder.summary_builder.create_tab(builder.workbook)
    SensitivityTabBuilder().create_tab(builder.workbook)
    return builder


def test_the_workbook_status_cell_carries_the_alert():
    builder = _builder()
    builder._apply_publication_labels({
        "publication_allowed": True,
        "confidence_alert": {**ALERT, "text": SENTENCE, "label": "low confidence"},
        "range_low": 89.59, "range_high": 98.34,
    })
    summary = builder.workbook["Summary"]
    assert summary["B24"].value == "PUBLISHABLE (low confidence: far from analyst consensus)"
    assert summary["G24"].value == SENTENCE
    # The value is published under its ordinary label, not as an audit midpoint.
    assert "not published" not in str(summary["A26"].value)


def test_a_corroborated_workbook_is_plainly_publishable():
    builder = _builder()
    builder._apply_publication_labels({"publication_allowed": True})
    summary = builder.workbook["Summary"]
    assert summary["B24"].value == "PUBLISHABLE"
    assert summary["G24"].value == "The deterministic valuation publication checks passed."


def test_a_bank_workbook_relabelled_after_evaluation_carries_the_alert():
    import openpyxl

    from src.agents.fm.financial_model_builder import FinancialModelBuilder

    class Logger:
        def info(self, _message):
            pass

    builder = FinancialModelBuilder("JPM", Logger())
    builder.workbook = openpyxl.Workbook()
    builder.workbook.active.title = "Summary"
    builder.workbook.create_sheet("Bank Valuation")
    builder.summary_builder = type("Summary", (), {"currency": "USD",
                                                    "comps_included_in_blend": False})()
    builder._apply_publication_labels({
        "publication_allowed": True,
        "confidence_alert": {**ALERT, "text": SENTENCE, "label": "low confidence"},
    })
    bank = builder.workbook["Bank Valuation"]
    assert bank["B22"].value == "PUBLISHABLE (low confidence: far from analyst consensus)"
    assert bank["B23"].value == SENTENCE
    assert builder.workbook["Summary"]["A26"].value == "Bank fair value per share"


def test_a_range_only_workbook_never_shows_an_alert():
    builder = _builder()
    builder._apply_publication_labels({
        "publication_allowed": False,
        "withheld_reason": "The valuation methods span more than 1.8x.",
        "confidence_alert": {**ALERT, "text": SENTENCE},
        "range_low": 150.0, "range_high": 280.0,
    })
    summary = builder.workbook["Summary"]
    assert summary["B24"].value == "SCENARIO RANGE ONLY"
    assert summary["G24"].value == "The valuation methods span more than 1.8x."


# --------------------------------------------------------------------------
# The chat answer
# --------------------------------------------------------------------------

def _flagged_state(alert=ALERT, **metrics):
    return SimpleNamespace(
        financial_data=SimpleNamespace(key_metrics={"basic_info": {"currency": "USD"}}),
        financial_model=SimpleNamespace(valuation_metrics={
            "fair_value": 380.0, "current_price": 751.66, "upside_vs_market": -0.494,
            "point_estimate_withheld": False,
            **({"confidence_alert": alert} if alert else {}), **metrics,
        }),
        report=SimpleNamespace(content=(
            "## Investment Rating: SELL\n\n**12-Month Price Target**: USD 380.00\n"
        )),
        news_analysis=None,
    )


def _runner(state=None):
    runner = supervisor_agent.SupervisorWorkflowRunner.__new__(
        supervisor_agent.SupervisorWorkflowRunner)
    runner.ticker = "META"
    runner.state = state or _flagged_state()
    return runner


def test_the_fixed_statement_states_the_rating_and_then_the_alert():
    runner = _runner()
    out = runner._guard_user_answer("The report rating is BUY.")
    assert out == (
        "META audited report headline: investment rating SELL; model fair value "
        "$380.00 USD; 12-month price target USD 380.00.\n\n" + SENTENCE
    )
    assert runner.guard_template == out and runner.guard_kind == "published:SELL"


def test_prose_that_gives_vynn_s_view_gets_the_alert_it_left_out():
    runner = _runner()
    prose = "The report rates META a SELL with a model fair value of $380.00."
    assert runner._guard_user_answer(prose) == f"{SENTENCE}\n\n{prose}"
    # Quoting the fair value alone is giving VYNN's view too.
    figure = "On our numbers the shares are worth about $380.00 each."
    assert runner._guard_user_answer(figure) == f"{SENTENCE}\n\n{figure}"
    # Not a template: the prose stands, with the alert in front of it.
    assert runner.guard_template is None


def test_a_broad_valuation_answer_always_carries_the_alert():
    runner = _runner()
    prose = "Meta trades at 24x forward earnings and revenue grew 22% last quarter."
    assert runner._guard_user_answer(prose, require_full_benchmark=True) == f"{SENTENCE}\n\n{prose}"


@pytest.mark.parametrize("stated", [
    "Low confidence: the model is 49% below the price while analysts sit 15% above it.",
    "VYNN rates it SELL, but confidence is low: analysts' mean target is above the price.",
    "A low-confidence SELL: the Street's targets point the other way.",
])
def test_prose_that_already_states_the_alert_is_not_given_it_twice(stated):
    runner = _runner()
    prose = f"The report rates META a SELL with a model fair value of $380.00. {stated}"
    assert runner._guard_user_answer(prose, require_full_benchmark=True) == prose


def test_an_answer_about_something_else_is_left_alone():
    runner = _runner()
    for prose in (
        "META closed at $751.66, up 1.2% on the day.",
        "Insiders did not sell shares last quarter, and the buyback continues.",
    ):
        assert runner._guard_user_answer(prose) == prose


def test_a_corroborated_view_gets_no_alert():
    runner = _runner(_flagged_state(alert=None))
    prose = "The report rates META a SELL with a model fair value of $380.00."
    assert runner._guard_user_answer(prose) == prose
    assert runner._guard_user_answer(prose, require_full_benchmark=True) == prose
    template = runner._safe_published_valuation_answer({"rating": "SELL"})
    assert "Low confidence" not in template


def test_a_range_only_run_never_gets_an_alert_even_if_one_is_left_in_state():
    runner = _runner(_flagged_state(
        point_estimate_withheld=True, perpetual_price=300.0, exit_multiple_price=460.0,
        publication_withheld_reason="The valuation methods span more than 1.8x.",
    ))
    out = runner._guard_user_answer("META is a SELL at $380.00.")
    assert "VYNN's answer here is a range, not a single fair value" in out
    assert "Low confidence" not in out and "380.00" not in out


def test_prose_in_chinese_gets_the_alert_in_chinese():
    runner = _runner()
    prose = "报告评级为卖出，模型公允价值为 380.00 美元，低于当前股价。"
    assert runner._guard_user_answer(prose) == f"{alert_sentence_zh(ALERT)}\n\n{prose}"
    assert alert_sentence_zh(ALERT).startswith("置信度低：VYNN 的公允价值低于市场价 49%")


@pytest.mark.parametrize("prose", [
    "レポートの評価は売りで、モデルの公正価値は 380.00 ドルです。",          # Japanese
    "보고서 등급은 매도이며 모델 공정가치는 380.00 달러입니다.",              # Korean
    "El informe califica META como SELL con un valor razonable de $380.00.",   # Spanish
])
def test_other_languages_get_the_english_alert_never_the_chinese_one(prose):
    # Japanese shares its ideographs with Chinese; the Chinese sentence in a
    # Japanese answer would be the wrong language, not a translation.
    runner = _runner()
    assert runner._guard_user_answer(prose) == f"{SENTENCE}\n\n{prose}"


def test_a_chinese_question_reads_the_alert_in_its_first_line():
    ag = _chat_agent("Meta的报告评级是什么？分析一下", _flagged_state(), "META")
    _, out = _analysis(ag, _Provider("ignored"), "报告评级为买入，公允价值 500.00 美元。")
    first, rest = out.split("\n\n", 1)
    assert first == (
        "META：VYNN 评级为卖出。下方英文部分列出模型公允价值、12个月目标价与分析师基准。"
        + alert_sentence_zh(ALERT)
    )
    assert rest.startswith("META audited report headline: investment rating SELL")
    assert SENTENCE in rest


def test_follow_up_turns_remember_the_alert():
    ag = _chat_agent("what is META worth?", _flagged_state(), "META")
    valuation = ag._collect_analysis_results()["valuation"]
    assert valuation["fair_value"] == 380.0
    assert valuation["confidence_alert"] == SENTENCE
    plain = _chat_agent("what is META worth?", _flagged_state(alert=None), "META")
    assert "confidence_alert" not in plain._collect_analysis_results()["valuation"]


# --------------------------------------------------------------------------
# What the prose model and the findings card are handed
# --------------------------------------------------------------------------

def test_tool_results_hand_the_prose_model_the_sentence():
    from src.agents.tools.analysis_tools import (
        CONFIDENCE_ALERT_NOTE, NO_SINGLE_VALUE_NOTE, _alert_payload,
    )
    metrics = {"confidence_alert": ALERT}
    assert _alert_payload(metrics, withheld=False) == {
        "confidence_alert": SENTENCE, "rating_confidence": "low",
    }
    assert _alert_payload(metrics, withheld=True) == {}
    assert _alert_payload({}, withheld=False) == {}
    assert _alert_payload(None, withheld=False) == {}
    assert _alert_payload({"confidence_alert": "junk"}, withheld=False) == {}
    assert "LOW confidence" in CONFIDENCE_ALERT_NOTE
    assert "never call the result not rated, unrated or withheld" in NO_SINGLE_VALUE_NOTE
    # The prose model repeats field names, so none of them says "range" for a
    # result that may be one estimate, and none says "withheld".
    for field in ("`no_single_fair_value`", "`no_single_fair_value_reason`",
                  "`valuation_support_shape`"):
        assert field in NO_SINGLE_VALUE_NOTE
    assert "a scenario estimate, not a fair value" in NO_SINGLE_VALUE_NOTE
    assert "range_only" not in NO_SINGLE_VALUE_NOTE


def test_the_bounded_headline_never_hands_over_the_retired_label():
    from src.agents.tools.analysis_tools import _bounded_report_headline
    published = "### Investment Rating: SELL\n**12-Month Price Target**: $380.00\n"
    assert _bounded_report_headline(published, point_estimate_withheld=False) == {
        "rating": "SELL", "price_target_12m": "$380.00",
    }
    for content in (published, "### Investment Rating: NOT RATED\n",
                    "### Investment View: Range Only\n", None):
        assert _bounded_report_headline(content, point_estimate_withheld=True) == {
            "valuation_view": "no_single_fair_value",
        }
    # A range-only report states no rating, in the old wording or the new.
    assert "rating" not in _bounded_report_headline(
        "### Investment View: Range Only\n", point_estimate_withheld=False)


def test_the_findings_card_flags_a_fair_value_the_street_does_not_back():
    result = {"status": "ok", "currency": "USD", "fair_value": 380.0,
              "upside_vs_market": -0.494, "confidence_alert": SENTENCE}
    flagged = extract_findings("build_model", result)[0]
    assert flagged["label"] == "Fair value" and flagged["sub"].endswith(" vs market · low confidence")
    plain = extract_findings("build_model", {k: v for k, v in result.items() if k != "confidence_alert"})[0]
    assert plain["sub"].endswith(" vs market") and "low confidence" not in plain["sub"]


# --------------------------------------------------------------------------
# The retired labels are not accepted from a prose model either
# --------------------------------------------------------------------------

@pytest.mark.parametrize("text", [
    "TEST is NOT RATED.", "The stock is not-rated here.", "It remains unrated.",
    "The point estimate was withheld.", "VYNN is withholding a rating for now.",
    "We withhold the fair value.",
])
def test_retired_labels_are_recognised(text):
    assert supervisor_agent._RETIRED_LABELS.search(text)


@pytest.mark.parametrize("text", [
    "Dividends are subject to a 15% withholding tax for ADR holders.",
    "Moody's rated the notes A1 and the company is rated investment grade.",
    "Management withheld guidance for the second half.",
])
def test_ordinary_prose_is_not_mistaken_for_a_retired_label(text):
    assert not supervisor_agent._RETIRED_LABELS.search(text)


@pytest.mark.parametrize("text", [
    "VYNN's answer here is a range, not a single fair value.",
    "There is no single fair value for this run.",
    "VYNN's answer here is the evidence, not a fair value.",
    "The model gives a range rather than one fair value.",
    "VYNN has no fair value to state here.",
])
def test_the_position_statements_are_recognised(text):
    assert supervisor_agent._STATES_NO_SINGLE_VALUE.search(text)


# --------------------------------------------------------------------------
# A range is only called a range: one surviving method, or none, has its own words
# --------------------------------------------------------------------------

def _range_only_report(perpetual, exit_value, *, comps=None, reason="One method failed."):
    from src.report_agent import apply_valuation_override
    data = apply_valuation_override(_valuation_data(comps=None), {
        "valuation_method": "dcf", "perpetual_price": perpetual,
        "exit_multiple_price": exit_value, "comps_price": comps,
        "dispersion_band": "single-method", "point_estimate_withheld": True,
        "publication_withheld_reason": reason,
    })
    data["news"] = {
        "summary": {"overall_sentiment": "neutral"},
        "catalysts": [], "risks": [], "freshness": {"status": "unavailable"},
    }
    return data


@pytest.mark.parametrize("perpetual,exit_value,view,status,table,closing,thesis,commentary,warning", [
    (188.30, 744.86, "Range Only", "Range only",
     "Range only: the evidence does not support one value",
     "VYNN shows the range above instead of a single fair value, rating or price target, because",
     "VYNN gives the range and the evidence here, not a buy or sell call.",
     "the range is the answer", "RANGE ONLY: this run supports a valuation range"),
    (264.73, -17.33, "Scenario Estimate Only", "Scenario estimate only",
     "Scenario estimate only: the evidence does not support a fair value",
     "VYNN shows the scenario estimate above, not a fair value, rating or price target, because",
     "VYNN gives the scenario estimate and the evidence here, not a buy or sell call.",
     "the scenario estimate is the answer: calling it a fair value",
     "SCENARIO ESTIMATE ONLY: this run supports one scenario estimate"),
    (-4.10, -17.33, "No Fair Value", "None",
     "None: no method produced a usable value",
     "VYNN states no fair value, rating or price target here, because",
     "VYNN gives the evidence here, not a buy or sell call.",
     "there is no fair value to state", "NO FAIR VALUE: no valuation method produced a usable value"),
])
def test_every_report_section_follows_what_the_methods_support(
        perpetual, exit_value, view, status, table, closing, thesis, commentary, warning):
    from src.report_agent import (
        generate_executive_summary, generate_section_investment_thesis,
        generate_section_valuation, valuation_publication_status,
    )
    data = _range_only_report(perpetual, exit_value)
    reliability = data["valuation"]["reliability"]
    assert reliability["point_estimate_withheld"] is True
    assert reliability["warning"].startswith(warning)

    status_block = valuation_publication_status(data)
    assert f"**Single Fair Value**: {status}" in status_block
    assert closing in status_block

    valuation_text, _ = generate_section_valuation(data, _must_not_call)
    assert f"| **Single Fair Value** | **{table}** |" in valuation_text
    assert commentary in valuation_text

    thesis_text, _ = generate_section_investment_thesis(data, _must_not_call)
    assert thesis in thesis_text
    assert "**Why no single fair value**: One method failed." in thesis_text

    summary, _ = generate_executive_summary(
        {"recommendation": "_Section unavailable (recommendation)._"}, data, None)
    assert summary.splitlines()[0] == f"**Investment View**: {view}"

    everything = "\n".join((status_block, valuation_text, thesis_text, summary))
    assert "NOT RATED" not in everything and "ithheld" not in everything
    if view != "Range Only":
        # Nothing calls one estimate, or no estimate, a range.
        for sentence in ("shows a range", "the range is the answer", "Range only", "Range Only",
                         "a valuation range", "the range and the evidence"):
            assert sentence not in everything, sentence


def test_the_chat_answer_follows_the_same_three_shapes():
    def answer(**metrics):
        runner = _runner(SimpleNamespace(
            financial_data=SimpleNamespace(key_metrics={"basic_info": {"currency": "USD"}}, raw_data={}),
            financial_model=SimpleNamespace(assumptions={}, valuation_metrics={
                "point_estimate_withheld": True, "current_price": 100.0,
                "publication_withheld_reason": "At least one valuation method failed with a non-positive value (exit_multiple_dcf).",
                **metrics,
            }),
            report=None, news_analysis=None,
        ))
        return runner._safe_withheld_valuation_answer()

    both = answer(perpetual_price=60.0, exit_multiple_price=140.0,
                  publication_withheld_reason="The valuation methods span more than 1.8x.")
    assert "VYNN's answer here is a range, not a single fair value, because" in both
    assert "The supported valuation-method range is $60.00–$140.00 USD." in both

    one = answer(perpetual_price=60.0, exit_multiple_price=-5.0)
    assert "VYNN's answer here is a scenario estimate, not a fair value, because" in one
    assert "The supported DCF scenario estimate is $60.00 USD." in one

    none = answer(perpetual_price=-3.0, exit_multiple_price=-5.0)
    assert "VYNN has no fair value to state here, because" in none
    assert "No supported positive valuation-method estimate is available." in none
    for text in (one, none):
        assert "is a range" not in text
    for text in (both, one, none):
        assert "NOT RATED" not in text and "withheld" not in text.lower()
        # Every one of them satisfies the guard's own test for stating the position.
        assert supervisor_agent._STATES_NO_SINGLE_VALUE.search(text)


def test_one_estimate_has_no_midpoint_in_any_model_view(engine, tmp_path, monkeypatch):
    from src.summary_evidence import model_view_summary, supported_valuation_span

    one = model_view_summary(span=supported_valuation_span([26.05]), current_price=372.11)
    assert one["headline"] == "the modeled cash flows support a value below the market (-93%)."
    both = model_view_summary(span=supported_valuation_span([150.0, 190.0]), current_price=300.0)
    assert both["headline"] == (
        "the modeled cash flows support a value below the market (-43% at the midpoint).")
    unpriced = model_view_summary(span=supported_valuation_span([26.05]), current_price=None)
    assert unpriced["headline"] == (
        "the modeled cash flows support one scenario estimate, but no market "
        "price was available to compare it with.")
    assert "a valuation range, but no market price" in model_view_summary(
        span=supported_valuation_span([150.0, 190.0]), current_price=None)["headline"]

    # The report's fallback recommendation states the same view.
    numbers = _unrated_numbers(264.73, 264.73)
    numbers["current_price"] = 300.0
    text = engine._evidence_safe_recommendation(numbers, {"evidence": []}, {})
    assert "**Model View**: The modeled cash flows support a value near the market (-12%)" in text
    assert "at the midpoint" not in text
    numbers = _unrated_numbers(150.0, 290.0)
    numbers["current_price"] = 300.0
    assert "(-27% at the midpoint)" in engine._evidence_safe_recommendation(
        numbers, {"evidence": []}, {})
