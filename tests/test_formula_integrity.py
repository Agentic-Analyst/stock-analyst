import openpyxl

from src.agents.fm.formula_evaluator import (
    FormulaEvaluator,
    formula_integrity,
    model_integrity,
)


def _evaluate(workbook, order):
    evaluator = FormulaEvaluator(workbook)
    evaluator.TAB_ORDER = order
    return evaluator.evaluate_all_tabs()


def test_cross_sheet_range_with_repeated_sheet_endpoint_evaluates():
    workbook = openpyxl.Workbook()
    source = workbook.active
    source.title = "Valuation (DCF)"
    for column, value in enumerate((.91, .84, .77, .70, .64), start=2):
        source.cell(row=17, column=column, value=value)
    summary = workbook.create_sheet("Summary")
    summary["B43"] = "=MAX('Valuation (DCF)'!$B$17:'Valuation (DCF)'!$F$17)<=1"

    result = _evaluate(workbook, ["Valuation (DCF)", "Summary"])

    assert result["Summary"]["cells"]["(43, 2)"] is True
    assert formula_integrity(result)["status"] == "ready"


def test_unsupported_function_is_an_integrity_error_not_a_placeholder():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "=TOTALLY_UNKNOWN(1)"

    result = _evaluate(workbook, ["Sheet"])
    manifest = formula_integrity(result)

    assert manifest["status"] == "error"
    assert manifest["issue_count"] == 1
    assert "Unsupported Excel function" in manifest["issues"][0]["error"]


def test_failed_reference_cannot_silently_turn_into_zero():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "=MissingSheet!A1+1"

    result = _evaluate(workbook, ["Sheet"])
    value = result["Sheet"]["cells"]["(1, 1)"]

    assert isinstance(value, dict)
    assert "unknown worksheet" in value["error"]
    assert formula_integrity(result)["issue_count"] == 1


def test_nested_function_failure_cannot_silently_turn_into_zero():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "=1+TOTALLY_UNKNOWN(2)"

    result = _evaluate(workbook, ["Sheet"])
    value = result["Sheet"]["cells"]["(1, 1)"]

    assert isinstance(value, dict)
    assert "nested function" in value["error"]
    assert formula_integrity(result)["status"] == "error"


def test_unguarded_division_by_zero_is_an_integrity_error():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "=10/0"

    result = _evaluate(workbook, ["Sheet"])

    assert "Division by zero" in result["Sheet"]["cells"]["(1, 1)"]["error"]
    assert formula_integrity(result)["status"] == "error"


def test_iferror_can_explicitly_handle_division_by_zero():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = '=IFERROR(10/0,"")'

    result = _evaluate(workbook, ["Sheet"])

    assert result["Sheet"]["cells"]["(1, 1)"] == ""
    assert formula_integrity(result)["status"] == "ready"


def test_missing_match_is_not_row_zero_and_can_be_handled_by_iferror():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "present"
    workbook.active["B1"] = '=MATCH("missing",A1:A1,0)'
    workbook.active["C1"] = '=IFERROR(MATCH("missing",A1:A1,0),99)'

    result = _evaluate(workbook, ["Sheet"])

    assert "MATCH could not find" in result["Sheet"]["cells"]["(1, 2)"]["error"]
    assert result["Sheet"]["cells"]["(1, 3)"] == 99


def test_referenced_text_cannot_be_executed_as_python_arithmetic():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "__import__('os').system('echo unsafe')"
    workbook.active["B1"] = "=A1+1"

    result = _evaluate(workbook, ["Sheet"])

    error = result["Sheet"]["cells"]["(1, 2)"]["error"]
    assert "Cannot evaluate arithmetic" in error
    assert formula_integrity(result)["status"] == "error"


def test_excel_boolean_arithmetic_used_for_method_counts_is_supported():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = 10
    workbook.active["A2"] = -2
    workbook.active["B1"] = "=(A1>0)+(A2>0)"

    result = _evaluate(workbook, ["Sheet"])

    assert result["Sheet"]["cells"]["(1, 2)"] == 1.0
    assert formula_integrity(result)["status"] == "ready"


def _coherent_model_results():
    """Small evaluated model whose accounting/valuation identities all tie."""
    historical = {
        "(1, 6)": 2025,
        "(3, 6)": 100.0, "(4, 6)": 60.0, "(5, 6)": 40.0,
        "(6, 6)": 5.0, "(7, 6)": 15.0, "(8, 6)": 20.0,
        "(18, 6)": 5.0, "(19, 6)": 25.0,
        "(23, 6)": 30.0, "(24, 6)": -10.0,
        "(25, 6)": 20.0, "(26, 6)": 20.0, "(27, 6)": 0.0,
        "(30, 6)": 10.0, "(31, 6)": 2.0, "(35, 6)": 5.0,
        "(37, 6)": -7.0, "(38, 6)": 50.0, "(44, 6)": 20.0,
    }
    projections = {}
    prior_ppe = 50.0
    prior_nwc = 7.0
    for col in range(2, 7):
        revenue = 100.0 + 10.0 * (col - 1)
        cost, rd, sga = revenue * .60, revenue * .05, revenue * .15
        gross = revenue - cost
        ebit = gross - rd - sga
        tax, nopat = ebit * .20, ebit * .80
        da, capex = revenue * .05, -revenue * .06
        ar, inventory, ap = revenue * .10, revenue * .05, revenue * .08
        nwc = ar + inventory - ap
        delta_nwc = nwc - prior_nwc
        fcf = nopat + da + capex - delta_nwc
        ending_ppe = prior_ppe - capex - da
        values = {
            3: revenue, 4: cost, 5: gross, 7: rd, 8: sga, 9: ebit,
            10: tax, 11: nopat, 12: da, 13: capex, 14: ar,
            15: inventory, 16: ap, 17: nwc, 18: delta_nwc, 19: fcf,
            43: prior_ppe, 44: -capex, 45: -da, 46: ending_ppe,
        }
        projections.update({f"({row}, {col})": value for row, value in values.items()})
        prior_ppe, prior_nwc = ending_ppe, nwc

    dcf = {
        "(6, 2)": .10, "(8, 2)": .20, "(9, 2)": .04, "(10, 2)": .80,
        "(11, 2)": .20, "(12, 2)": .088,
        "(26, 2)": 100.0, "(30, 2)": 10.0, "(31, 2)": 5.0,
        "(32, 2)": 2.0, "(36, 2)": 10.0,
    }
    explicit_pv = 0.0
    for col in range(2, 12):
        fcf = projections.get(f"(19, {col})", 20.0)
        discount = 1.0 / (1.088 ** (col - 1))
        dcf[f"(16, {col})"] = fcf
        dcf[f"(17, {col})"] = discount
        dcf[f"(18, {col})"] = fcf * discount
        explicit_pv += fcf * discount
    dcf["(19, 2)"] = explicit_pv
    dcf["(27, 2)"] = explicit_pv + 100.0
    dcf["(33, 2)"] = dcf["(27, 2)"] + 10.0 - 5.0 + 2.0
    dcf["(37, 2)"] = dcf["(33, 2)"] / 10.0

    exit_tab = {
        "(3, 2)": 8.0, "(4, 2)": .20, "(10, 2)": explicit_pv, "(12, 2)": 20.0,
        "(13, 2)": 8.0, "(14, 2)": 160.0, "(15, 2)": 80.0,
        "(17, 2)": explicit_pv + 80.0, "(19, 2)": 10.0,
        "(20, 2)": 5.0, "(21, 2)": 2.0, "(24, 2)": 10.0,
    }
    exit_tab["(22, 2)"] = exit_tab["(17, 2)"] + 10.0 - 5.0 + 2.0
    exit_tab["(25, 2)"] = exit_tab["(22, 2)"] / 10.0
    for col in range(2, 12):
        exit_tab[f"(7, {col})"] = dcf[f"(16, {col})"]
        exit_tab[f"(8, {col})"] = dcf[f"(17, {col})"]
        exit_tab[f"(9, {col})"] = dcf[f"(18, {col})"]

    summary = {
        "(8, 2)": 10.0, "(9, 2)": 20.0, "(10, 2)": 200.0,
        "(13, 2)": dcf["(33, 2)"], "(14, 2)": 10.0,
        "(15, 2)": 5.0, "(16, 2)": 2.0, "(17, 2)": dcf["(27, 2)"],
        "(18, 2)": dcf["(37, 2)"], "(20, 2)": exit_tab["(17, 2)"],
        "(21, 2)": exit_tab["(22, 2)"], "(22, 2)": exit_tab["(25, 2)"],
        "(29, 2)": 193.0,
    }
    summary["(26, 2)"] = (summary["(18, 2)"] + summary["(22, 2)"]) / 2
    summary["(27, 2)"] = summary["(26, 2)"] / summary["(9, 2)"] - 1
    summary.update({f"({row}, 2)": True for row in range(41, 47)})
    return {
        "Model_Inputs": {"cells": {"(11, 2)": .20}},
        "Assumptions": {"cells": {"(20, 2)": .20}},
        "Historical": {"cells": historical},
        "Projections": {"cells": projections},
        "Valuation (DCF)": {"cells": dcf},
        "Valuation (Exit Multiple)": {"cells": exit_tab},
        "Summary": {"cells": summary},
    }


def test_semantic_model_integrity_accepts_a_tied_model():
    manifest = model_integrity(_coherent_model_results())

    assert manifest == {
        "status": "ready", "issue_count": 0,
        "issues": [], "issues_truncated": False,
    }


def test_semantic_model_integrity_catches_wrong_cell_references_that_evaluate():
    result = _coherent_model_results()
    result["Historical"]["cells"]["(44, 6)"] = 120.0
    result["Projections"]["cells"]["(43, 2)"] = 20.0
    result["Summary"]["cells"]["(18, 2)"] += 1.0

    manifest = model_integrity(result)
    checks = {issue["check"] for issue in manifest["issues"]}

    assert manifest["status"] == "error"
    assert "historical_ebit_tieout" in checks
    assert "ppe_beginning_balance" in checks
    assert "summary_perpetual_value_per_share" in checks


def test_semantic_bank_integrity_accepts_an_absent_peer_method():
    results = {
        "Bank Valuation": {"cells": {
            "(5, 2)": 100.0,
            "(6, 2)": .15,
            "(7, 2)": .09,
            "(8, 2)": .03,
            "(10, 2)": 2.0,
            "(11, 2)": 2.0,
            "(12, 2)": 200.0,
            "(14, 2)": None,
            "(15, 2)": None,
            "(18, 2)": 200.0,
            "(19, 2)": 200.0,
            "(20, 2)": 200.0,
        }},
        "Summary": {"cells": {
            "(4, 2)": .09,
            "(5, 2)": .03,
            "(6, 2)": .15,
            "(7, 2)": 2.0,
            "(18, 2)": 200.0,
            # FormulaEvaluator represents a reference to an empty cell as 0.
            "(22, 2)": 0,
            "(26, 2)": 200.0,
        }},
    }

    manifest = model_integrity(results)

    assert manifest["status"] == "ready"
    assert manifest["issues"] == []


def test_semantic_bank_integrity_rejects_a_phantom_peer_value():
    results = {
        "Bank Valuation": {"cells": {
            "(5, 2)": 100.0, "(6, 2)": .15, "(7, 2)": .09,
            "(8, 2)": .03, "(10, 2)": 2.0, "(11, 2)": 2.0,
            "(12, 2)": 200.0, "(14, 2)": None, "(15, 2)": None,
            "(18, 2)": 200.0, "(19, 2)": 200.0, "(20, 2)": 200.0,
        }},
        "Summary": {"cells": {
            "(4, 2)": .09, "(5, 2)": .03, "(6, 2)": .15,
            "(7, 2)": 2.0, "(18, 2)": 200.0,
            "(22, 2)": 250.0, "(26, 2)": 200.0,
        }},
    }

    manifest = model_integrity(results)

    assert manifest["status"] == "error"
    assert manifest["issues"][0]["check"] == "bank_summary_peer_absent"
