import inspect

import openpyxl
import pytest

from src.agents.fm.financial_model_builder import FinancialModelBuilder
from src.agents.fm.tabs.tab_assumptions import (
    AssumptionsTabBuilder,
    infer_assumptions_with_llm,
    source_grounded_bank_assumption_seed,
    source_grounded_assumption_seed,
)


def _payload(*, operating_income=-25.0):
    return {
        "financial_statements": {
            "income_statement": {
                "2025-12-31": {
                    "Total Revenue": 100.0,
                    "Gross Profit": 40.0,
                    "Operating Income": operating_income,
                    "EBITDA": operating_income + 5.0,
                    "Cost Of Revenue": 60.0,
                },
                "2024-12-31": {
                    "Total Revenue": 80.0,
                    "Gross Profit": 32.0,
                    "Operating Income": operating_income - 2.0,
                    "EBITDA": operating_income + 3.0,
                    "Cost Of Revenue": 48.0,
                },
            },
            "balance_sheet": {
                "2025-12-31": {
                    "Accounts Receivable": 10.0,
                    "Inventory": 6.0,
                    "Accounts Payable": 9.0,
                }
            },
        },
        "modeling_metrics": {
            "historical_growth_rates": {"revenue_growth": {"cagr_3y": 0.20}}
        },
    }


def test_numerical_assumption_seed_is_deterministic_and_source_grounded():
    first = source_grounded_assumption_seed(_payload())
    second = source_grounded_assumption_seed(_payload())

    assert first == second
    assert first["revenue_growth_rates"] == pytest.approx(
        [0.20, 0.165, 0.13, 0.095, 0.06]
    )
    assert first["gross_margins"] == pytest.approx([0.40] * 5)
    assert first["ebitda_margins"] == pytest.approx([-0.20] * 5)
    assert first["operating_margins"] == pytest.approx([-0.25] * 5)
    assert first["assumption_seed_source"] == "three_year_revenue_cagr"


def test_loss_making_company_never_receives_generic_software_margins():
    assumptions = source_grounded_assumption_seed(_payload(operating_income=-60.0))

    assert assumptions["operating_margins"] == pytest.approx([-0.60] * 5)
    assert assumptions["gross_margins"] == pytest.approx([0.40] * 5)
    assert assumptions["operating_margins"] != pytest.approx([0.31] * 5)


def test_legacy_inference_entry_point_no_longer_calls_a_language_model():
    payload = _payload()

    assert infer_assumptions_with_llm(payload) == source_grounded_assumption_seed(payload)
    builder_source = inspect.getsource(FinancialModelBuilder.build_model)
    assert "source_grounded_assumption_seed" in builder_source
    assert "infer_assumptions_with_llm" not in builder_source


def test_missing_core_statement_inputs_fail_closed():
    payload = _payload()
    del payload["financial_statements"]["income_statement"]["2025-12-31"][
        "Operating Income"
    ]

    with pytest.raises(ValueError, match="revenue and operating income"):
        source_grounded_assumption_seed(payload)


def test_bank_seed_does_not_invent_or_require_industrial_operating_margin():
    payload = _payload()
    del payload["financial_statements"]["income_statement"]["2025-12-31"][
        "Operating Income"
    ]

    assumptions = source_grounded_bank_assumption_seed(payload)

    assert assumptions["assumption_seed_source"] == "not_applicable_to_bank_valuation"
    assert assumptions["operating_margins"] == []
    assert assumptions["revenue_growth_rates"] == []


def test_workbook_exposes_model_inputs_and_contains_no_generic_defaults():
    workbook = openpyxl.Workbook()
    assumptions = source_grounded_assumption_seed(_payload(operating_income=20.0))
    AssumptionsTabBuilder(assumptions).create_tab(workbook)

    assert "Model_Inputs" in workbook.sheetnames
    assert "LLM_Inferred" not in workbook.sheetnames
    assert workbook["Model_Inputs"]["B7"].value == pytest.approx(0.20)
    assert workbook["Assumptions"]["C7"].value == (
        '=IF(IFERROR(Model_Inputs!B4,"")<>"",Model_Inputs!B4,B7)'
    )
