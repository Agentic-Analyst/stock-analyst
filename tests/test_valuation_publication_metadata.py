from datetime import datetime, timedelta, timezone

import pytest

from src.valuation_publication_metadata import (
    build_publication_metadata,
    fail_closed_publication_metadata,
)


def _computed(perpetual=50.0, exit_value=60.0, price=100.0):
    cells = {
        "(1, 1)": "Current Market Price", "(1, 2)": price,
        "(2, 1)": "Value per Share (Perpetual DCF)", "(2, 2)": perpetual,
        "(3, 1)": "Value per Share (Exit Multiple DCF)", "(3, 2)": exit_value,
        "(4, 1)": "DCF Fair Value (Per-Share)",
        "(4, 2)": (perpetual + exit_value) / 2,
        # These are the canonical Summary-tab rows consumed by the report and
        # persistence boundary.  The display labels above intentionally do
        # not substitute for the workbook schema.
        "(18, 2)": perpetual,
        "(22, 2)": exit_value,
        "(26, 2)": (perpetual + exit_value) / 2,
        "(27, 2)": ((perpetual + exit_value) / 2) / price - 1,
    }
    return {
        "Summary": {"cells": cells, "metadata": {"total_cells": len(cells)}},
        "Valuation (Exit Multiple)": {
            "cells": {"(25, 2)": exit_value},
            "metadata": {"total_cells": 1},
        },
    }


def _financials():
    latest = (datetime.now(timezone.utc) - timedelta(days=90)).date().isoformat()
    prior = (datetime.now(timezone.utc) - timedelta(days=455)).date().isoformat()
    return {
        "company_data": {
            "basic_info": {"quote_type": "EQUITY", "currency": "USD"},
            "market_data": {"current_price": 100.0, "market_cap": 8e9},
            "forward_guidance": {
                "target_mean_price": 112.0,
                "number_of_analyst_opinions": 18,
            },
        },
        "financial_statements": {
            "income_statement": {
                latest: {"Total Revenue": 100},
                prior: {"Total Revenue": 90},
            },
            "balance_sheet": {
                latest: {"Total Assets": 200},
                prior: {"Total Assets": 180},
            },
            "cash_flow": {
                latest: {"Operating Cash Flow": 20, "Free Cash Flow": 15},
                prior: {"Operating Cash Flow": 18, "Free Cash Flow": 13},
            },
        },
    }


def test_model_artifact_carries_the_same_street_conflict_boundary_as_report():
    metadata = build_publication_metadata(_computed(), _financials())

    assert metadata["status"] == "ready"
    assert metadata["point_estimate_withheld"] is True
    assert metadata["valuation_conclusion"] == "inconclusive"
    assert metadata["canonical_fair_value"] is None
    assert metadata["model_value_for_audit"] == 55.0
    assert "18-analyst target benchmark" in metadata["withheld_reason"]


def test_model_artifact_aligns_street_revenue_to_rolling_forecast_clock():
    computed = _computed(perpetual=90.0, exit_value=110.0, price=100.0)
    computed["Projections"] = {
        "cells": {
            "(3, 2)": 110.0,
            "(3, 3)": 150.0,
            "(19, 2)": 20.0,
            "(19, 3)": 22.0,
        },
        "metadata": {"total_cells": 4},
    }
    financials = _financials()
    financials["external_expectations"] = {
        "forward_estimates": [
            {
                "period": "0y", "revenue": 100.0,
                "revenue_analyst_count": 12,
            },
            {
                "period": "+1y", "revenue": 140.0,
                "revenue_analyst_count": 10,
            },
        ],
    }
    assumptions = {
        "forecast_basis": {
            "basis": "rolling_twelve_months",
            "period_end": "2026-06-30",
            "fiscal_year_progress": 0.25,
            "base_revenue": 90.0,
        },
    }

    metadata = build_publication_metadata(
        computed, financials, assumptions=assumptions,
    )

    assert "near-term operating case is not reconciled" not in (
        metadata.get("withheld_reason") or ""
    )
    assert "model revenue" not in (metadata.get("withheld_reason") or "")


def test_extreme_market_gap_marks_artifact_as_modeled_cash_flow_scope_only():
    computed = _computed(perpetual=10.0, exit_value=20.0, price=100.0)
    computed["Valuation (DCF)"] = {
        "cells": {"(27, 2)": 10.0},
        "metadata": {"total_cells": 1},
    }
    computed["Summary"]["cells"]["(51, 2)"] = 430.7

    metadata = build_publication_metadata(computed, _financials())

    assert metadata["point_estimate_withheld"] is True
    assert metadata["market_implied_fcf_path_vs_model"] == pytest.approx(42.07)
    assert metadata["model_scope"] == "modeled_operating_cash_flow_path_only"
    assert "not a comprehensive company value" in metadata["model_scope_warning"]


def test_publication_metadata_failure_is_explicit_and_fail_closed():
    metadata = fail_closed_publication_metadata(RuntimeError("private detail"))

    assert metadata["status"] == "error"
    assert metadata["point_estimate_withheld"] is True
    assert metadata["canonical_fair_value"] is None
    assert "private detail" not in repr(metadata)


def test_downloadable_summary_relabels_withheld_midpoint_as_diagnostic():
    import openpyxl

    from src.agents.fm.financial_model_builder import FinancialModelBuilder
    from src.agents.fm.tabs.tab_sensitivity import SensitivityTabBuilder
    from src.agents.fm.tabs.tab_summary import SummaryTabBuilder

    class Logger:
        def info(self, _message):
            pass

    builder = FinancialModelBuilder("AAPL", Logger())
    builder.workbook = openpyxl.Workbook()
    builder.workbook.remove(builder.workbook.active)
    builder.summary_builder = SummaryTabBuilder(currency="USD")
    builder.summary_builder.create_tab(builder.workbook)
    SensitivityTabBuilder().create_tab(builder.workbook)
    builder._apply_publication_labels({
        "publication_allowed": False,
        "withheld_reason": "Independent evidence conflicts.",
        "range_low": 150.0,
        "range_high": 180.0,
    })

    summary = builder.workbook["Summary"]
    assert summary["B24"].value == "WITHHELD — scenario range only"
    assert summary["B25"].value == "150.00 to 180.00 USD/share"
    assert summary["A26"].value == "DCF scenario midpoint (not published)"
    assert summary["A27"].value == "Audit midpoint vs market (diagnostic)"
    assert summary["B26"].value.startswith("=IF(")
    assert builder.workbook["Sensitivity"]["A34"].value == (
        "DCF scenario midpoint (not published)"
    )
    assert builder.workbook["Sensitivity"]["A36"].value == (
        "Audit midpoint vs market (diagnostic)"
    )


def test_zero_width_display_range_is_labeled_as_one_scenario_estimate():
    import openpyxl

    from src.agents.fm.financial_model_builder import FinancialModelBuilder
    from src.agents.fm.tabs.tab_summary import SummaryTabBuilder

    builder = FinancialModelBuilder("TSM", type("Logger", (), {})())
    builder.workbook = openpyxl.Workbook()
    builder.workbook.remove(builder.workbook.active)
    builder.summary_builder = SummaryTabBuilder(currency="TWD")
    builder.summary_builder.create_tab(builder.workbook)
    builder._apply_publication_labels({
        "publication_allowed": False,
        "withheld_reason": "Independent corroboration is unavailable.",
        "range_low": 13245.003,
        "range_high": 13245.004,
    })

    summary = builder.workbook["Summary"]
    assert summary["B24"].value == "WITHHELD — scenario estimate only"
    assert summary["A25"].value == "Supported DCF scenario estimate"
    assert summary["B25"].value == "13,245.00 TWD/share"


def test_post_evaluation_labels_are_synchronized_to_computed_sidecar():
    import openpyxl

    from src.agents.fm.financial_model_builder import FinancialModelBuilder
    from src.agents.fm.tabs.tab_sensitivity import SensitivityTabBuilder
    from src.agents.fm.tabs.tab_summary import SummaryTabBuilder

    builder = FinancialModelBuilder("AAPL", type("Logger", (), {})())
    builder.workbook = openpyxl.Workbook()
    builder.workbook.active.title = "Summary"
    builder.summary_builder = SummaryTabBuilder(currency="USD")
    builder.workbook.remove(builder.workbook["Summary"])
    builder.summary_builder.create_tab(builder.workbook)
    SensitivityTabBuilder().create_tab(builder.workbook)
    results = {
        "Summary": {"cells": {"(26, 1)": "DCF Fair Value (Per-Share)"}},
        "Sensitivity": {"cells": {"(34, 1)": "DCF Scenario Midpoint"}},
    }

    builder._apply_publication_labels({
        "publication_allowed": False,
        "withheld_reason": "Independent evidence conflicts.",
        "range_low": 150.0,
        "range_high": 180.0,
    })
    builder._sync_publication_labels_into_results(results)

    cells = results["Summary"]["cells"]
    assert cells["(24, 2)"].startswith("WITHHELD")
    assert cells["(26, 1)"] == "DCF scenario midpoint (not published)"
    assert cells["(27, 1)"] == "Audit midpoint vs market (diagnostic)"
    assert results["Sensitivity"]["cells"]["(34, 1)"].endswith(
        "(not published)"
    )
    assert results["Sensitivity"]["cells"]["(36, 1)"] == (
        "Audit midpoint vs market (diagnostic)"
    )


def test_downloadable_bank_summary_uses_publication_specific_headline():
    import openpyxl

    from src.agents.fm.financial_model_builder import FinancialModelBuilder

    class Logger:
        def info(self, _message):
            pass

    builder = FinancialModelBuilder("JPM", Logger())
    builder.workbook = openpyxl.Workbook()
    builder.workbook.active.title = "Summary"
    builder.workbook.create_sheet("Bank Valuation")
    builder.summary_builder = type("Summary", (), {"currency": "USD",
                                                    "comps_included_in_blend": False})()
    builder._apply_publication_labels({
        "publication_allowed": False,
        "withheld_reason": "External evidence conflicts.",
    })

    assert builder.workbook["Summary"]["A26"].value.endswith("(not published)")
    assert builder.workbook["Bank Valuation"]["B22"].value.startswith("WITHHELD")
    assert builder.workbook["Bank Valuation"]["B23"].value == "External evidence conflicts."


def test_bank_publication_metadata_never_calls_itself_a_corporate_dcf(monkeypatch):
    computed = _computed(perpetual=90.0, exit_value=110.0, price=100.0)
    computed["Bank Valuation"] = {
        "cells": {
            "(4, 2)": 100.0,
            "(5, 2)": 50.0,
            "(6, 2)": 0.15,
            "(7, 2)": 0.10,
            "(8, 2)": 0.025,
            "(10, 2)": 1.6666666667,
            "(11, 2)": 1.6666666667,
            "(12, 2)": 83.33,
            "(18, 2)": 83.33,
            "(19, 2)": 83.33,
            "(20, 2)": 83.33,
            "(21, 2)": -0.1667,
        },
        "metadata": {"total_cells": 12},
    }
    monkeypatch.setattr(
        "src.agents.fm.bank_valuation.build_bank_valuation_override",
        lambda *args, **kwargs: {
            "valuation_method": "justified_pb_roe",
            "fair_value": 83.33,
            "intrinsic_fair_value": 83.33,
            "forward_consensus_fair_value": None,
            "peer_fair_value": None,
            "current_price": 100.0,
            "upside_vs_market": -0.1667,
            "bank_inputs": {
                "bvps": 50.0,
                "roe": 0.15,
                "cost_of_equity": 0.10,
                "terminal_growth": 0.025,
                "justified_pb": 1.6666,
                "intrinsic_fair_value": 83.33,
            },
        },
    )

    metadata = build_publication_metadata(computed, _financials())

    assert metadata["model_scope"] == "bank_balance_sheet_valuation"


def test_nonfinite_model_value_is_a_completed_fail_closed_policy_decision():
    metadata = build_publication_metadata(
        _computed(perpetual=float("nan"), exit_value=float("nan")),
        _financials(),
    )

    assert metadata["status"] == "ready"
    assert metadata["point_estimate_withheld"] is True
    assert metadata["canonical_fair_value"] is None
    assert metadata["model_value_for_audit"] is None
    assert "No finite positive" in metadata["withheld_reason"]


def test_negative_scenario_value_remains_auditable_but_never_publishable():
    metadata = build_publication_metadata(
        _computed(perpetual=-90.0, exit_value=-30.0),
        _financials(),
    )

    assert metadata["status"] == "ready"
    assert metadata["point_estimate_withheld"] is True
    assert metadata["canonical_fair_value"] is None
    assert metadata["model_value_for_audit"] == -60.0
    assert metadata["range_low"] is None
    assert metadata["range_high"] is None
    assert metadata["method_values_for_audit"] == {
        "perpetual_dcf": -90.0,
        "exit_multiple_dcf": -30.0,
    }
    assert metadata["failed_method_values"] == metadata["method_values_for_audit"]


def test_failed_method_is_auditable_but_not_a_supported_range_endpoint():
    metadata = build_publication_metadata(
        _computed(perpetual=100.0, exit_value=-30.0),
        _financials(),
    )

    assert metadata["point_estimate_withheld"] is True
    assert metadata["range_low"] == 100.0
    assert metadata["range_high"] == 100.0
    assert metadata["method_values_for_audit"] == {
        "perpetual_dcf": 100.0,
        "exit_multiple_dcf": -30.0,
    }
    assert metadata["failed_method_values"] == {"exit_multiple_dcf": -30.0}


def test_well_covered_street_earnings_conflict_blocks_model_only_publication():
    financials = _financials()
    financials["company_data"]["market_data"].update({
        "shares_outstanding_basic": 10.0,
    })
    financials["analyst_data"] = {
        "revenue_estimates": {
            "0y": {"avg": 100.0, "numberOfAnalysts": 12},
            "+1y": {"avg": 110.0, "numberOfAnalysts": 12},
        },
        "earnings_estimates": {
            "0y": {"avg": 1.0, "numberOfAnalysts": 12},
            "+1y": {"avg": 1.1, "numberOfAnalysts": 12},
        },
    }
    computed = _computed(perpetual=95.0, exit_value=105.0)
    computed["Projections"] = {"cells": {
        "(3, 2)": 100.0, "(3, 3)": 110.0,
        "(11, 2)": 35.0, "(11, 3)": 38.5,
    }}

    metadata = build_publication_metadata(computed, financials)

    assert metadata["point_estimate_withheld"] is True
    assert metadata["canonical_fair_value"] is None
    assert "profitability case is not reconciled" in metadata["withheld_reason"]
    assert "12 EPS analysts" in metadata["withheld_reason"]


def test_saved_bank_publication_metadata_can_rehydrate_report_override():
    from src.report_agent import valuation_override_from_publication_metadata

    computed = {"_vynn": {"valuation_publication": {
        "status": "ready",
        "valuation_method": "justified_pb_roe",
        "point_estimate_withheld": True,
        "withheld_reason": "External evidence conflicts.",
        "valuation_confidence": "tight",
        "model_value_for_audit": 288.93,
        "valuation_method_inputs": {
            "book_value_per_share": 133.0,
            "return_on_equity": 0.169,
            "intrinsic_fair_value": 281.23,
            "forward_consensus_fair_value": 312.14,
            "forward_consensus_return_on_equity": .18,
            "forward_consensus_justified_price_to_book": 2.35,
            "peer_fair_value": 296.63,
            "peer_observation_count": 6,
            "input_boundary_triggered": True,
            "book_value_cross_check_failed": True,
            "book_value_provider_gap": .20,
        },
    }}}

    override = valuation_override_from_publication_metadata(computed)

    assert override["fair_value"] == 288.93
    assert override["intrinsic_fair_value"] == 281.23
    assert override["forward_consensus_fair_value"] == 312.14
    assert override["peer_fair_value"] == 296.63
    assert override["bank_inputs"]["bvps"] == 133.0
    assert override["bank_inputs"]["input_boundary_triggered"] is True
    assert override["bank_inputs"]["book_value_cross_check_failed"] is True
    assert override["point_estimate_withheld"] is True


def test_bank_metadata_publishes_only_the_appropriate_method_and_inputs():
    financials = _financials()
    financials["company_data"].update({
        "basic_info": {
            "quote_type": "EQUITY", "currency": "USD",
            "sector": "Financial Services", "industry": "Banks - Diversified",
        },
        "valuation_metrics": {"book_value": 100.0},
        "growth_profitability": {"return_on_equity": .16},
        "market_data": {"current_price": 150.0, "market_cap": 1e11},
        "capital_structure": {"beta": 1.0},
        "forward_guidance": {
            "target_mean_price": 210.0,
            "number_of_analyst_opinions": 18,
        },
    })
    financials["modeling_metrics"] = {"financial_ratios": {"financial_profile": {
        "interest_income_to_revenue": 1.2,
    }}}

    metadata = build_publication_metadata(
        _computed(perpetual=10.0, exit_value=20.0, price=150.0),
        financials,
        assumptions={
            "terminal_growth_rate": .025,
            "capm": {
                "risk_free_rate": .04,
                "equity_risk_premium_total": .05,
                "beta": 1.0,
            },
        },
    )

    assert metadata["valuation_method"] == "justified_pb_roe"
    assert metadata["valuation_confidence"] == "single-method"
    # The Yahoo/legacy target has no provider as-of date. It remains a visible
    # external benchmark, but cannot positively validate a 38% single-method
    # point estimate.
    assert metadata["canonical_fair_value"] is None
    assert metadata["point_estimate_withheld"] is True
    assert metadata["model_value_for_audit"] == 207.69
    assert metadata["valuation_method_inputs"]["book_value_per_share"] == 100.0
    assert metadata["valuation_method_inputs"]["justified_price_to_book"] == pytest.approx(
        2.0769230769)
    assert metadata["comps_included_in_blended_value"] is False
