"""Summary quality-control cells must evaluate to booleans, not error objects."""

import logging

import openpyxl

from src.agents.fm.formula_evaluator import FormulaEvaluator
from src.agents.fm.tabs.tab_summary import SummaryTabBuilder


def _evaluate_qa(*, perpetual_factors, exit_factors):
    workbook = openpyxl.Workbook()
    perpetual = workbook.active
    perpetual.title = "Valuation (DCF)"
    perpetual["B10"] = 0.8
    perpetual["B11"] = 0.2
    perpetual["B12"] = 0.09
    perpetual["B23"] = 0.025
    for column, value in zip("BCDEFGHIJK", perpetual_factors):
        perpetual[f"{column}17"] = value

    exit_multiple = workbook.create_sheet("Valuation (Exit Multiple)")
    for column, value in zip("BCDEFGHIJK", exit_factors):
        exit_multiple[f"{column}8"] = value

    sensitivity = workbook.create_sheet("Sensitivity")
    sensitivity["B2"] = "No"

    summary = workbook.create_sheet("Summary")
    summary["B8"] = 100.0
    summary["B9"] = 10.0
    SummaryTabBuilder()._setup_qa_flags(summary)

    evaluator = FormulaEvaluator(workbook)
    evaluator.set_logger(logging.getLogger("test_summary_qa"))
    return evaluator.evaluate_all_tabs()["Summary"]["cells"]


def test_valid_discount_factor_ranges_return_true_without_errors():
    cells = _evaluate_qa(
        perpetual_factors=(0.92, 0.84, 0.77, 0.71, 0.65, 0.60, 0.55, 0.51, 0.47, 0.43),
        exit_factors=(0.91, 0.83, 0.76, 0.70, 0.64, 0.59, 0.54, 0.50, 0.46, 0.42),
    )

    assert [cells[f"({row}, 2)"] for row in range(41, 47)] == [True] * 6


def test_discount_factor_above_one_fails_the_relevant_check():
    cells = _evaluate_qa(
        perpetual_factors=(0.92, 0.84, 0.77, 0.71, 0.65, 0.60, 0.55, 0.51, 0.47, 1.01),
        exit_factors=(0.91, 0.83, 0.76, 0.70, 0.64, 0.59, 0.54, 0.50, 0.46, 0.42),
    )

    assert cells["(43, 2)"] is False
    assert cells["(44, 2)"] is True
