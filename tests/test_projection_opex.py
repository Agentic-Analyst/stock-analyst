"""Operating-expense projection must survive incomplete line-item disclosure."""

import logging

import openpyxl
import pytest

from src.agents.fm.formula_evaluator import FormulaEvaluator
from src.agents.fm.tabs.tab_projections import ProjectionsTabBuilder


def _evaluate_projection(*, historical_rd, historical_sga, operating_margin=0.05):
    workbook = openpyxl.Workbook()
    historical = workbook.active
    historical.title = "Historical"
    historical["F3"] = 100.0
    historical["F6"] = historical_rd
    historical["F7"] = historical_sga
    historical["F18"] = 2.0
    historical["F24"] = -3.0
    historical["F32"] = 0.0
    historical["F33"] = 0.0
    historical["F34"] = 0.0
    historical["F38"] = 50.0

    assumptions = workbook.create_sheet("Assumptions")
    assumptions["B2"] = 2025
    assumptions["B20"] = 0.20
    for column in "CDEFG":
        assumptions[f"{column}7"] = 0.10
        assumptions[f"{column}9"] = 0.25
        assumptions[f"{column}13"] = 0.0
        assumptions[f"{column}14"] = 0.0
        assumptions[f"{column}15"] = 0.0

    inferred = workbook.create_sheet("Model_Inputs")
    for column in "BCDEF":
        inferred[f"{column}7"] = operating_margin

    ProjectionsTabBuilder().create_tab(workbook)
    evaluator = FormulaEvaluator(workbook)
    evaluator.set_logger(logging.getLogger("test_projection_opex"))
    return evaluator.evaluate_all_tabs()["Projections"]["cells"]


def test_missing_expense_breakout_does_not_delete_total_opex():
    cells = _evaluate_projection(historical_rd=0.0, historical_sga=0.0)

    # Revenue 110, gross profit 27.5, target EBIT 5.5: all 22 of otherwise
    # unclassified opex is retained in SG&A instead of disappearing.
    assert cells["(7, 2)"] == 0.0
    assert cells["(8, 2)"] == pytest.approx(22.0)
    assert cells["(9, 2)"] == pytest.approx(5.5)
    assert cells["(9, 2)"] / cells["(3, 2)"] == pytest.approx(0.05)


def test_disclosed_expenses_keep_their_historical_split_and_target_margin():
    cells = _evaluate_projection(historical_rd=3.0, historical_sga=7.0)

    assert cells["(7, 2)"] == pytest.approx(6.6)
    assert cells["(8, 2)"] == pytest.approx(15.4)
    assert cells["(9, 2)"] / cells["(3, 2)"] == pytest.approx(0.05)


def test_ppe_rollforward_starts_from_net_ppe_not_historical_free_cash_flow():
    cells = _evaluate_projection(historical_rd=3.0, historical_sga=7.0)

    assert cells["(43, 2)"] == pytest.approx(50.0)


def test_reinvestment_rate_uses_net_capex_and_working_capital_with_correct_signs():
    cells = _evaluate_projection(historical_rd=3.0, historical_sga=7.0)

    # Revenue 110; gross capex 3.3; D&A 2.2; delta NWC 0; NOPAT 4.4.
    assert cells["(50, 2)"] == pytest.approx((3.3 - 2.2) / 4.4)
