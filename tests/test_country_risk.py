"""
Country risk premiums come from Damodaran's published table, not a hand-typed one.

Run:  python -m pytest tests/test_country_risk.py -q
"""

import datetime as dt
import io
import os
import sys
import time

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "src"))

from src.agents.fm import netcache
from src.agents.fm import country_risk as cr


def _workbook(n_filler=110, as_of=dt.datetime(2026, 1, 1)):
    """A sheet laid out like ctryprem.xlsx, including the rows that must be skipped."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Explanation and FAQ"
    ws = wb.create_sheet("ERPs by country")
    ws.append(["Country and Equity Risk Premiums"])
    ws.append(["Date of update:", as_of, "(Updated Feb 16 for sovereign ratings"])
    ws.append(["Enter the current risk premium for a mature equity market", None, None, None, 0.0423])
    ws.append(["Enter the current risk premium for the US =", None, None, None, 0.0446])
    ws.append([None])
    ws.append(["Country", "Africa", "Moody's rating", "Rating-based Default Spread",
               "Total Equity Risk Premium", "Country Risk Premium", "Sovereign CDS, net of Swiss CDS",
               "Total Equity Risk Premium2", "Country Risk Premium3"])
    rows = [
        ("United States", "North America", "Aa1", 0.002334, 0.0446, 0.002334),
        ("India", "Asia", "Baa3", 0.018678, 0.070754, 0.028454),
        ("Germany", "Western Europe", "Aaa", 0, 0.0423, 0),
        ("Korea", "Asia", "Aa2", 0.004195, 0.048691, 0.006391),
        ("Czech Republic", "Eastern Europe & Russia", "Aa3", 0.005094, 0.05006, 0.00776),
        ("Japan", "Asia", "A1", 0.005993, 0.051429, 0.009129),
        ("United Kingdom", "Western Europe", "Aa3", 0.005094, 0.05006, 0.00776),
    ]
    for r in rows:
        ws.append(list(r))
    for i in range(n_filler):
        ws.append([f"Country {i}", "Region", "Baa2", 0.016181, 0.06695, 0.02465])
    # Frontier and malformed rows: numbers where the rating should be.
    ws.append(["Frontier Markets (no sovereign ratings)"])
    ws.append(["Algeria", 67, 0.100577, 0.058277, 0.038255])
    ws.append(["Russia", 70.25, 0.081253, 0.038953, 0.02557])
    ws.append([None, "Baa2", 161.8])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestParsing:
    def test_reads_the_date_the_mature_erp_and_every_rated_row(self):
        t = cr.parse_workbook(_workbook())
        assert t["as_of"] == "2026-01-01"
        assert t["mature_erp"] == pytest.approx(0.0423)
        assert t["countries"]["India"] == {"rating": "Baa3", "default_spread": pytest.approx(0.018678),
                                           "crp": pytest.approx(0.028454)}
        assert t["countries"]["Germany"]["crp"] == 0.0

    def test_frontier_and_malformed_rows_are_skipped_not_guessed(self):
        t = cr.parse_workbook(_workbook())
        for name in ("Algeria", "Russia", "Frontier Markets (no sovereign ratings)", "Country"):
            assert name not in t["countries"]

    def test_a_sheet_that_is_not_the_table_is_refused(self):
        with pytest.raises(ValueError):
            cr.parse_workbook(_workbook(n_filler=3))
        import openpyxl
        wb = openpyxl.Workbook()
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(ValueError):
            cr.parse_workbook(buf.getvalue())


class TestLoading:
    def test_the_live_file_is_used_and_cached(self, monkeypatch):
        monkeypatch.setattr(netcache, "http_get", lambda url, timeout=None: _workbook())
        t = cr.load_table()
        assert t["source"] == "Damodaran January 2026 update"
        assert "snapshot" not in t["source"]
        assert os.path.exists(os.path.join(netcache.cache_dir(), "ctryprem.json"))

        cr._reset_memo()
        monkeypatch.setattr(netcache, "http_get", lambda *a, **k: (_ for _ in ()).throw(ConnectionError()))
        assert cr.load_table()["source"] == "Damodaran January 2026 update"   # from disk

    def test_offline_falls_to_the_snapshot_and_says_so(self):
        t = cr.load_table()
        assert "snapshot" in t["source"]
        assert len(t["countries"]) > 150
        assert t["countries"]["India"]["rating"] == "Baa3"

    def test_a_corrupt_download_does_not_replace_the_snapshot(self, monkeypatch):
        monkeypatch.setattr(netcache, "http_get", lambda url, timeout=None: b"<html>maintenance</html>")
        assert "snapshot" in cr.load_table()["source"]

    def test_an_expired_disk_copy_beats_the_embedded_snapshot(self, monkeypatch):
        """The July update fetched eight days ago is newer than the January snapshot in the code."""
        monkeypatch.setattr(netcache, "http_get", lambda url, timeout=None: _workbook(as_of=dt.datetime(2026, 7, 1)))
        assert cr.load_table()["source"] == "Damodaran July 2026 update"
        cr._reset_memo()
        path = os.path.join(netcache.cache_dir(), "ctryprem.json")
        old = time.time() - 8 * 86400
        os.utime(path, (old, old))
        monkeypatch.setattr(netcache, "http_get", lambda *a, **k: (_ for _ in ()).throw(ConnectionError()))
        t = cr.load_table()
        assert t["source"] == "Damodaran July 2026 update, cached"
        assert t["as_of"] == "2026-07-01"

    @pytest.mark.parametrize("countries", [
        {"India": {"rating": "Baa3"}},                                   # no spreads: would KeyError
        {f"C{i}": {"rating": "Baa3", "default_spread": 0.01, "crp": 0.02} for i in range(101)} | {"X": "not a dict"},
        {f"C{i}": {"rating": "Baa3", "default_spread": 9.0, "crp": 0.02} for i in range(101)},   # 900%
    ])
    def test_a_malformed_cache_file_is_a_miss_not_a_crash(self, countries):
        import json
        with open(os.path.join(netcache.cache_dir(), "ctryprem.json"), "w") as fh:
            json.dump({"as_of": "2026-01-01", "source": "Damodaran January 2026 update", "countries": countries}, fh)
        t = cr.load_table()
        assert "snapshot" in t["source"]
        assert cr.sovereign_default_spread("INR")[0] > 0


class TestLookup:
    def test_yahoo_s_names_are_understood(self):
        assert cr.country_entry("South Korea")["name"] == "Korea"
        assert cr.country_entry("Czechia")["name"] == "Czech Republic"
        assert cr.country_entry("india")["name"] == "India"
        # Damodaran writes "Jersey (States of)", "Congo (Republic of)"; Yahoo writes the short form.
        assert cr.country_entry("Jersey")["name"] == "Jersey (States of)"
        assert cr.country_entry("Guernsey")["rating"]
        assert cr.country_entry("Atlantis") is None
        assert cr.country_entry(None) is None

    def test_premium_labels_name_the_source_the_rating_and_the_date(self):
        value, label = cr.country_risk_premium("India")
        assert 0.02 <= value <= 0.04
        assert "Damodaran" in label and "2026" in label and "Baa3" in label

    def test_aaa_sovereigns_carry_no_premium(self):
        value, label = cr.country_risk_premium("Germany")
        assert value == 0.0 and "Aaa" in label and "no country premium" in label

    def test_the_united_states_is_rated_aa1_and_carries_a_small_one(self):
        value, label = cr.country_risk_premium("United States")
        assert 0.001 < value < 0.005 and "Aa1" in label

    def test_unknown_country_is_not_guessed(self):
        value, label = cr.country_risk_premium("Atlantis")
        assert value == 0.0 and "no country premium on record" in label

    def test_override(self, monkeypatch):
        monkeypatch.setenv("CRP_INDIA", "0.035")
        value, label = cr.country_risk_premium("India")
        assert value == pytest.approx(0.035) and "override" in label


class TestDefaultSpread:
    def test_the_currency_s_sovereign_decides(self):
        value, label = cr.sovereign_default_spread("INR")
        assert 0.015 < value < 0.025 and "India" in label and "Baa3" in label

    def test_euro_is_the_bund_which_is_aaa(self):
        assert cr.sovereign_default_spread("EUR") == (0.0, "")

    def test_pence_is_sterling(self):
        value, label = cr.sovereign_default_spread("GBp")
        assert value > 0 and "United Kingdom" in label

    def test_unknown_currency_has_none(self):
        assert cr.sovereign_default_spread("XXX") == (0.0, "")
